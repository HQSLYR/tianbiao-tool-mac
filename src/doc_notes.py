# -*- coding: utf-8 -*-
"""把用户给的文档（图片/表格/Word/文本）整理成日志或笔记。

- 种类：日志（按时间戳逐条记录） / 笔记（分章节整理）
- 话语：简单（摘要式短句） / 具体（原文完整内容）
- 落地：Word(.docx，带图片) / 记事本(.txt，图片以说明文字标注)

图片用 Windows 自带 OCR（WinRT，PowerShell 调用）识别文字，
系统缺中文 OCR 组件时自动退到用户语言/英文，再不行记一条说明，不报错。
来料里的图片（图片文件本身、Word 内嵌图）会一并写进日志 Word。
"""
import hashlib
import os
import re
import subprocess
import tempfile
import time
import zipfile

import xlsx_fill

# 内容条数上限，防止几万行的表把笔记撑爆
_MAX_ROWS_DETAIL = 200
_MAX_ROWS_SIMPLE = 30
_MAX_CHARS_SIMPLE = 40

_OCR_PS = r'''
$ErrorActionPreference = 'Stop'
[Console]::OutputEncoding = [System.Text.Encoding]::UTF8
Add-Type -AssemblyName System.Runtime.WindowsRuntime
$asTaskGeneric = ([System.WindowsRuntimeSystemExtensions].GetMethods() | Where-Object { $_.Name -eq 'AsTask' -and $_.GetParameters().Count -eq 1 -and $_.GetParameters()[0].ParameterType.Name -eq 'IAsyncOperation`1' })[0]
function Await($WinRtTask, $ResultType) {
  $asTask = $asTaskGeneric.MakeGenericMethod($ResultType)
  $netTask = $asTask.Invoke($null, @($WinRtTask))
  $netTask.Wait(-1) | Out-Null
  $netTask.Result
}
[Windows.Storage.StorageFile,Windows.Storage,ContentType=WindowsRuntime] | Out-Null
[Windows.Globalization.Language,Windows.Foundation,ContentType=WindowsRuntime] | Out-Null
[Windows.Media.Ocr.OcrEngine,Windows.Foundation,ContentType=WindowsRuntime] | Out-Null
[Windows.Graphics.Imaging.BitmapDecoder,Windows.Graphics,ContentType=WindowsRuntime] | Out-Null
$engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromLanguage([Windows.Globalization.Language]::new('zh-Hans-CN'))
if (-not $engine) { $engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromLanguage([Windows.Globalization.Language]::new('en-US')) }
if (-not $engine) { $engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages() }
if (-not $engine) { Write-Output 'OCR_UNAVAILABLE'; exit 0 }
$file = Await ([Windows.Storage.StorageFile]::GetFileFromPathAsync($args[0])) ([Windows.Storage.StorageFile])
$stream = Await ($file.OpenAsync([Windows.Storage.FileAccessMode]::Read)) ([Windows.Storage.Streams.IRandomAccessStream])
$decoder = Await ([Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($stream)) ([Windows.Graphics.Imaging.BitmapDecoder])
$bitmap = Await ($decoder.GetSoftwareBitmapAsync()) ([Windows.Graphics.Imaging.SoftwareBitmap])
$result = Await ($engine.RecognizeAsync($bitmap)) ([Windows.Media.Ocr.OcrResult])
Write-Output $result.Text
'''

_CN = '一二三四五六七八九十'


def _cn(i):
    return _CN[i - 1] if 1 <= i <= 10 else str(i)


def _short(t, n=_MAX_CHARS_SIMPLE):
    t = (t or '').strip()
    return t if len(t) <= n else t[:n] + '…'


def _first_sentence(t):
    t = (t or '').strip()
    for sep in ('。', '！', '？', '；', '\n'):
        if sep in t:
            return _short(t.split(sep)[0] + sep)
    return _short(t)


# ---------------- 各类文档取内容 ----------------
_IMG_MARK = re.compile(r'〔IMGS:(\d+)〕')   # 行内检索，兼容日志行的时间戳前缀
_MAX_DOCX_IMGS = 40          # 单个文档最多带出的内嵌图数
_MAX_IMG_BYTES = 20 << 20    # 单张超 20MB 的不随日志携带
_IMG_EXT = ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tiff', '.tif')


def _media_images(path, prefix, max_imgs=_MAX_DOCX_IMGS):
    """直接从压缩包(word/xl 的 media 目录)提取所有位图内嵌图，比只扫主文档
    related_parts 更全：页眉/页脚/文本框/表格里的图一并带走。跳过 emf/wmf 矢量图
    （python-docx 无法渲染）。返回临时目录里的图片路径列表。"""
    imgs = []
    try:
        z = zipfile.ZipFile(path)
    except Exception:
        return imgs
    tmpdir = tempfile.mkdtemp(prefix='notes_img_')
    try:
        names = [n for n in z.namelist()
                 if n.startswith(prefix)
                 and os.path.splitext(n)[1].lower() in _IMG_EXT]

        def keyf(n):   # 按文件名里的数字排序，尽量贴近文档出现顺序
            m = re.search(r'(\d+)', os.path.basename(n))
            return (int(m.group(1)) if m else 9999, n)
        seen = set()
        for n in sorted(names, key=keyf):
            try:
                blob = z.read(n)
            except Exception:
                continue
            if not blob or len(blob) > _MAX_IMG_BYTES:
                continue
            digest = hashlib.md5(blob).hexdigest()
            if digest in seen:      # 同一张图被多处引用只带一次
                continue
            seen.add(digest)
            ext = os.path.splitext(n)[1].lower().lstrip('.')
            if ext in ('jpeg', 'jfif'):
                ext = 'jpg'
            fp = os.path.join(tmpdir, 'img_%02d.%s' % (len(imgs) + 1, ext))
            try:
                with open(fp, 'wb') as f:
                    f.write(blob)
                imgs.append(fp)
            except OSError:
                continue
            if len(imgs) >= max_imgs:
                break
    finally:
        z.close()
    return imgs


def _docx_images(path):
    """提取 Word 内嵌图片（含页眉/页脚/正文），返回图片路径列表。"""
    return _media_images(path, 'word/media/')


def _ensure_docx(path):
    """旧版 .doc/.wps 自动转 .docx 再读；已是 .docx 原样返回。"""
    from word_to_table import to_docx
    return to_docx(path)


def _docx_lines(path, detailed):
    from docx import Document
    path = _ensure_docx(path)
    doc = Document(path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for tb in doc.tables:
        for row in tb.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                tables.append(' | '.join(cells))
    imgs = _docx_images(path)
    info = 'Word 文档，%d 个段落、%d 行表格内容' % (len(paras), len(tables))
    if imgs:
        info += '、%d 张内嵌图片' % len(imgs)
    if detailed:
        body = ['段落原文：'] + paras
        if tables:
            body += ['表格内容：'] + tables
    else:
        body = ['每段一句话摘要：'] + [_first_sentence(t) for t in paras[:_MAX_ROWS_SIMPLE]]
        if len(paras) > _MAX_ROWS_SIMPLE:
            body.append('……（其余 %d 段略）' % (len(paras) - _MAX_ROWS_SIMPLE))
    return info, body, imgs


def _xlsx_lines(path, detailed):
    header, rows = xlsx_fill.read_rows(path)
    cols = sorted(header, key=xlsx_fill.colkey)
    imgs = _media_images(path, 'xl/media/')
    info = '表格，%d 行数据、%d 列（%s…）' % (
        len(rows), len(cols), '、'.join(header[c] for c in cols[:3]))
    if imgs:
        info += '、%d 张内嵌图片' % len(imgs)
    body = []
    limit = _MAX_ROWS_DETAIL if detailed else _MAX_ROWS_SIMPLE
    for i, row in enumerate(rows[:limit], 2):
        vals = [(header.get(c, c), row[c]) for c in sorted(row, key=xlsx_fill.colkey)
                if row.get(c)]
        if detailed:
            body.append('第%d行：%s' % (i, '；'.join('%s=%s' % kv for kv in vals)))
        else:
            body.append('第%d行：%s' % (i, ' | '.join(v for _k, v in vals[:3])))
    if len(rows) > limit:
        body.append('……（其余 %d 行略）' % (len(rows) - limit))
    return info, body, imgs


def _txt_lines(path, detailed):
    raw = open(path, 'rb').read()
    for enc in ('utf-8-sig', 'utf-8', 'gbk'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode('gbk', errors='replace')
    lines = [t.strip() for t in text.splitlines() if t.strip()]
    info = '文本文件，%d 行有效内容' % len(lines)
    if detailed:
        body = lines[:2000]
    else:
        body = [_short(t) for t in lines[:_MAX_ROWS_SIMPLE]]
        if len(lines) > _MAX_ROWS_SIMPLE:
            body.append('……（其余 %d 行略）' % (len(lines) - _MAX_ROWS_SIMPLE))
    return info, body, []


def ocr_image(path):
    """Windows 自带 OCR 识别图片文字；失败返回 (False, 原因)。"""
    fd, ps = tempfile.mkstemp(suffix='.ps1')
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as f:
            f.write(_OCR_PS)
        p = subprocess.run(['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass',
                            '-File', ps, os.path.abspath(path)],
                           capture_output=True, timeout=120)
        out = p.stdout.decode('utf-8', errors='replace').strip()
    except Exception as e:
        return False, '调用系统 OCR 失败：%s' % e
    finally:
        try:
            os.remove(ps)
        except OSError:
            pass
    if p.returncode != 0:
        return False, 'OCR 执行出错：%s' % p.stderr.decode('utf-8', errors='replace')[:200]
    out = out.strip()
    if out == 'OCR_UNAVAILABLE':
        return False, '系统未装 OCR 组件（设置→时间和语言→语言→中文选项→OCR）'
    if not out:
        return False, 'OCR 没有识别到文字（图片可能无文字或太模糊）'
    return True, out


def _img_lines(path, detailed):
    size = os.path.getsize(path)
    ok, res = ocr_image(path)
    info = '图片（%.1f KB）' % (size / 1024.0)
    imgs = [os.path.abspath(path)] if size <= _MAX_IMG_BYTES else []
    if not ok:
        return info, ['识别说明：%s' % res], imgs
    lines = [t.strip() for t in res.splitlines() if t.strip()]
    info += '，OCR 识别出 %d 行文字' % len(lines)
    if detailed:
        body = ['图片文字全文：'] + lines
    else:
        body = ['图片文字摘要：'] + [_short(t) for t in lines[:10]]
        if len(lines) > 10:
            body.append('……（其余 %d 行略）' % (len(lines) - 10))
    return info, body, imgs


_EXT_MAP = {'.docx': _docx_lines, '.doc': _docx_lines, '.wps': _docx_lines,
            '.xlsx': _xlsx_lines, '.xls': _xlsx_lines,
            '.txt': _txt_lines, '.log': _txt_lines, '.md': _txt_lines,
            '.png': _img_lines, '.jpg': _img_lines, '.jpeg': _img_lines,
            '.bmp': _img_lines}


def extract(path, detailed):
    """返回 (类型说明, 内容行列表, 随日志携带的图片路径列表)；不支持的类型抛 ValueError。"""
    ext = os.path.splitext(path)[1].lower()
    fn = _EXT_MAP.get(ext)
    if not fn:
        raise ValueError('不支持的文件类型：%s（支持 图片/表格/Word/文本）' % ext)
    return fn(path, detailed)


# ---------------- 组装日志/笔记 ----------------
def build_notes(files, kind='笔记', detailed=True):
    """返回 (text, img_blocks)；img_blocks = {来源序号: [图片路径]}，
    text 内用〔IMGS:序号〕标记图片插入位（txt 落地时换成说明文字）。"""
    now = time.strftime('%Y-%m-%d %H:%M:%S')
    hm = time.strftime('%H:%M:%S')
    out = []
    img_blocks = {}
    if kind == '日志':
        out.append('==== 日志 ====')
        out.append('生成时间：%s' % now)
        out.append('')
        for i, path in enumerate(files, 1):
            info, body, imgs = extract(path, detailed)
            out.append('[%s] 来源%d：《%s》（%s）' % (
                hm, i, os.path.basename(path), info))
            for line in body:
                out.append('[%s]   %s' % (hm, line))
            if imgs:
                img_blocks[i] = imgs
                out.append('[%s] 〔IMGS:%d〕' % (hm, i))
            out.append('')
    else:
        out.append('笔记')
        out.append('生成时间：%s' % now)
        out.append('')
        for i, path in enumerate(files, 1):
            info, body, imgs = extract(path, detailed)
            out.append('%s、《%s》' % (_cn(i), os.path.basename(path)))
            out.append('（一）基本情况：%s' % info)
            out.append('（二）内容%s：' % ('（具体）' if detailed else '（简要）'))
            out.extend('    ' + line for line in body)
            if imgs:
                img_blocks[i] = imgs
                out.append('〔IMGS:%d〕' % i)
            out.append('')
    return '\n'.join(out), img_blocks


# ---------------- 落地 Word / 记事本 ----------------
def save_txt(text, path, img_blocks=None):
    img_blocks = img_blocks or {}
    lines = []
    for line in text.splitlines():
        m = _IMG_MARK.search(line)
        if m:
            n = len(img_blocks.get(int(m.group(1)), []))
            lines.append('    （本来源附 %d 张图片，图片见 Word 版日志）' % n)
        else:
            lines.append(line)
    with open(path, 'w', encoding='utf-8-sig') as f:   # 带 BOM，老记事本不乱码
        f.write('\n'.join(lines))


def save_docx(text, path, img_blocks=None):
    from docx import Document
    from docx.shared import Pt, Inches
    img_blocks = img_blocks or {}
    doc = Document()
    lines = text.splitlines()
    for i, line in enumerate(lines):
        m = _IMG_MARK.search(line)
        if m:
            for fp in img_blocks.get(int(m.group(1)), []):
                try:
                    doc.add_picture(fp, width=Inches(5.8))
                    doc.paragraphs[-1].alignment = 1   # 图片居中
                except Exception:
                    doc.add_paragraph('（图片插入失败：%s）' % os.path.basename(fp))
            continue
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.size = Pt(14 if i == 0 else 11)
        r.bold = (i == 0)
    doc.save(path)
