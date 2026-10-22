# -*- coding: utf-8 -*-
"""项目文书的只读资料解析与样式参考，不修改来源文件。"""
from collections import Counter
from copy import copy
from pathlib import Path
import re
import zipfile

ROLES = ('主要依据', '补充资料', '成品参考')
MAX_FILES = 10
MAX_BYTES = 20 * 1024 * 1024
MAX_CHARS = 300000


def check_package(path):
    """限制解压体积；拒绝宏文件及异常压缩包。"""
    with zipfile.ZipFile(path) as archive:
        if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError('Office 文件解压后超过100MB，请拆分资料。')
        if any('vbaproject' in name.lower() for name in archive.namelist()):
            raise ValueError('不接收包含宏的资料，请另存为不含宏的文件。')


def _word_blocks(path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    doc = Document(path)
    heading = ''
    pno = tno = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn('w:p'):
            pno += 1
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = p.style.name if p.style else ''
            if ('heading' in style.lower() or '标题' in style
                    or re.match(r'^[一二三四五六七八九十]+[、．.]', text)):
                heading = text[:100]
            yield '段落%d / %s' % (pno, heading), text, ''
        elif child.tag == qn('w:tbl'):
            tno += 1
            table = Table(child, doc)
            for rno, row in enumerate(table.rows, 1):
                seen = set()
                values = []
                for cno, cell in enumerate(row.cells, 1):
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    values.append('列%d=%s' % (cno, cell.text.strip()))
                if any(cell.text.strip() for cell in row.cells):
                    yield '表%d第%d行 / %s' % (tno, rno, heading), ' | '.join(values), ''
    if doc.element.body.xpath('.//w:drawing | .//w:pict | .//w:txbxContent'):
        yield '读取提示', '', '文档含图片、图形或文本框；首版未识别其内部内容，请将重要内容补为正文。'


def _excel_blocks(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=False, keep_links=False)
    cached = load_workbook(path, data_only=True, keep_links=False)
    try:
        for sheet in workbook.worksheets:
            if sheet.max_row * sheet.max_column > 500000:
                raise ValueError('工作表“%s”的使用区域过大，请清理空白格式或拆分。' % sheet.title)
            merged = ', '.join(str(r) for r in sheet.merged_cells.ranges)
            if merged:
                yield '%s / 合并区域' % sheet.title, merged, ''
            for row in sheet.iter_rows():
                values, warnings = [], []
                for cell in row:
                    if cell.value is None:
                        continue
                    value = cell.value
                    if cell.data_type == 'f':
                        value = cached[sheet.title][cell.coordinate].value
                        if value is None:
                            value = '【待补充：公式无缓存值】'
                            warnings.append('%s!%s 公式无缓存；请用Excel/WPS重算保存后再载入。'
                                            % (sheet.title, cell.coordinate))
                    values.append('%s=%s' % (cell.coordinate, value))
                if values:
                    yield '%s / 第%d行' % (sheet.title, row[0].row), ' | '.join(values), '\n'.join(warnings)
    finally:
        workbook.close()
        cached.close()


def read_sources(files):
    """files=[{'path':..., 'role':...}]；每段都有稳定的来源编号。"""
    if not files or len(files) > MAX_FILES:
        raise ValueError('请选择1至10份资料。')
    if not any(f.get('role') == '主要依据' for f in files):
        raise ValueError('至少需要一份“主要依据”。')
    result = {'blocks': [], 'warnings': [], 'files': [], 'characters': 0}
    for index, item in enumerate(files, 1):
        path = Path(item['path'])
        role = item.get('role', '主要依据')
        if role not in ROLES:
            raise ValueError('未知资料角色。')
        ext = path.suffix.lower()
        if ext not in ('.docx', '.xlsx', '.txt', '.md'):
            raise ValueError('暂不支持%s，请另存为docx、xlsx、txt或md。' % ext)
        if not path.is_file() or path.stat().st_size == 0:
            raise ValueError('资料为空或不存在：%s' % path.name)
        if path.stat().st_size > MAX_BYTES:
            raise ValueError('单份资料不能超过20MB：%s' % path.name)
        if ext in ('.docx', '.xlsx'):
            check_package(path)
            iterator = _word_blocks(path) if ext == '.docx' else _excel_blocks(path)
        else:
            raw = path.read_bytes()
            for encoding in ('utf-8-sig', 'gbk'):
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeDecodeError:
                    pass
            else:
                raise ValueError('不能识别文本编码，请另存为UTF-8：%s' % path.name)
            iterator = (('第%d行' % n, line.strip(), '')
                        for n, line in enumerate(text.splitlines(), 1))
        before = len(result['blocks'])
        try:
            for location, text, warning in iterator:
                if warning:
                    result['warnings'].append('%s：%s' % (path.name, warning))
                if not text.strip():
                    continue
                result['characters'] += len(text)
                if result['characters'] > MAX_CHARS:
                    raise ValueError('提取文字超过30万字符，请拆分后生成；未静默截断。')
                for offset in range(0, len(text), 2400):
                    ident = 'F%d-B%d' % (index, len(result['blocks']) + 1)
                    result['blocks'].append({
                        'id': ident, 'file': path.name, 'role': role,
                        'location': location + (' / 续段%d' % (offset // 2400 + 1) if offset else ''),
                        'text': text[offset:offset + 2400],
                    })
        except (ValueError, OSError):
            raise
        except Exception as exc:
            raise ValueError('不能完整解析%s（%s），请检查加密、损坏或另存格式。'
                             % (path.name, type(exc).__name__)) from exc
        if before == len(result['blocks']):
            raise ValueError('没有读到有效正文：%s；扫描件请先转换文字。' % path.name)
        result['files'].append({'name': path.name, 'role': role,
                                'blocks': len(result['blocks']) - before})
    return result


def batches(blocks, max_chars=6500):
    """所有来源块都会进入一次提取，不仅取文档开头。"""
    current, length = [], 0
    for block in blocks:
        size = len(block['text']) + len(block['location']) + 160
        if current and length + size > max_chars:
            yield current
            current, length = [], 0
        current.append(block)
        length += size
    if current:
        yield current


def verified_citations(citations, sources, allow_reference=False):
    known = {b['id']: b for b in sources['blocks']}
    valid = []
    for citation in citations if isinstance(citations, list) else []:
        if not isinstance(citation, dict):
            continue
        source_id = citation.get('source_id')
        if not isinstance(source_id, str):
            continue
        block = known.get(source_id)
        quote = citation.get('quote', '')
        if (block and isinstance(quote, str) and len(quote.strip()) >= 4
                and quote in block['text']
                and (allow_reference or block['role'] != '成品参考')):
            valid.append({'source_id': block['id'], 'quote': quote})
    return valid


def _font_data(paragraph):
    from docx.oxml.ns import qn
    run = next((r for r in paragraph.runs if r.text.strip()), None)
    style = paragraph.style
    fonts = [x for x in (run.font if run else None, style.font if style else None) if x]
    family = None
    for element in (run._r if run else None, style.element if style else None):
        if element is not None:
            found = element.xpath('.//w:rFonts')
            if found:
                family = found[0].get(qn('w:eastAsia')) or found[0].get(qn('w:ascii'))
                if family:
                    break
    size = next((f.size.pt for f in fonts if f.size), None)
    bold = next((f.bold for f in fonts if f.bold is not None), None)
    color = next((str(f.color.rgb) for f in fonts if f.color and f.color.type == 1), None)
    fmt = paragraph.paragraph_format
    inherited = style.paragraph_format if style else None
    data = {'font': family, 'size': size, 'bold': bold, 'color': color,
            'alignment': int(paragraph.alignment) if paragraph.alignment is not None else None}
    for key in ('space_before', 'space_after', 'first_line_indent'):
        value = getattr(fmt, key)
        if value is None and inherited is not None:
            value = getattr(inherited, key)
        if value is not None:
            data[key] = value.pt
    spacing = fmt.line_spacing
    if spacing is None and inherited is not None:
        spacing = inherited.line_spacing
    if isinstance(spacing, float):
        data['line_spacing'] = spacing
    elif spacing is not None and hasattr(spacing, 'pt'):
        data['line_spacing_points'] = spacing.pt
    return {k: v for k, v in data.items() if v is not None}


def word_style(path=None):
    if not path:
        return {}, []
    from docx import Document
    from docx.oxml.ns import qn
    if Path(path).suffix.lower() != '.docx':
        raise ValueError('Word样式参考必须是.docx。')
    check_package(path)
    doc = Document(path)
    section = doc.sections[0]
    result = {'page': {name: getattr(section, name).pt for name in
                       ('page_width', 'page_height', 'top_margin', 'bottom_margin',
                        'left_margin', 'right_margin') if getattr(section, name) is not None}}
    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    bodies = [p for p in paragraphs if len(p.text) > 40 and not p.style.name.lower().startswith('heading')]
    if bodies:
        common = Counter(str(_font_data(p)) for p in bodies).most_common(1)[0][0]
        result['body'] = next(_font_data(p) for p in bodies if str(_font_data(p)) == common)
    elif paragraphs:
        result['body'] = _font_data(paragraphs[-1])
    for level in (1, 2, 3):
        p = next((p for p in paragraphs if p.style.name.lower() == 'heading %d' % level), None)
        if p is not None:
            result['heading%d' % level] = _font_data(p)
    if doc.tables and doc.tables[0].rows:
        cell = doc.tables[0].rows[0].cells[0]
        if cell.paragraphs:
            result['table_header'] = _font_data(cell.paragraphs[0])
        shades = cell._tc.xpath('./w:tcPr/w:shd')
        if shades and re.fullmatch('[0-9a-fA-F]{6}', shades[0].get(qn('w:fill'), '')):
            result['header_fill'] = shades[0].get(qn('w:fill'))
        borders = cell._tc.xpath('./w:tcPr/w:tcBorders/*')
        if not borders:
            borders = doc.tables[0]._tbl.xpath('./w:tblPr/w:tblBorders/*')
        if borders:
            result['border'] = {key: borders[0].get(qn('w:' + key)) for key in ('val', 'sz', 'color')}
    warning = '只参考可识别的纸张、字体、段距和表格样式；不复制正文、徽标、签章及复杂浮动图形。'
    missing = [str(i) for i in (1, 2, 3) if 'heading%d' % i not in result]
    if missing:
        warning += '第%s级标题未识别，使用标准样式。' % '、'.join(missing)
    return result, [warning]


def excel_style(path=None):
    if not path:
        return {}, []
    from openpyxl import load_workbook
    if Path(path).suffix.lower() != '.xlsx':
        raise ValueError('Excel样式参考必须是.xlsx。')
    check_package(path)
    wb = load_workbook(path, keep_links=False)
    try:
        ws = next((s for s in wb.worksheets if s.sheet_state == 'visible'), wb.worksheets[0])
        cells = [c for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12),
                                         max_col=min(ws.max_column, 20)) for c in row if c.value is not None]
        header = next((c for c in cells if c.font.bold), cells[0] if cells else ws['A1'])
        body = next((c for c in cells if c.row > header.row), header)
        result = {'header': {}, 'body': {}}
        for label, cell in (('header', header), ('body', body)):
            for attr in ('font', 'fill', 'border', 'alignment'):
                result[label][attr] = copy(getattr(cell, attr))
        result['widths'] = {name: value.width for name, value in ws.column_dimensions.items()}
        result['orientation'] = ws.page_setup.orientation
        result['paperSize'] = ws.page_setup.paperSize
        result['margins'] = copy(ws.page_margins)
        return result, ['只参考首个可见工作表的表头、单元格、列宽与打印样式，不复制原数据或公式。']
    finally:
        wb.close()
