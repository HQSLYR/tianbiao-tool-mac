# -*- coding: utf-8 -*-
"""Word 扫描填表工具 - 图形界面版（扫 Word 内课程表，多行填写）

逻辑：
    ① 选要扫描的 Word；② 选 Excel 表格（表头决定要哪些列）。
    ③ 程序在 Word 里找含“课程编号/课程名称”列的课程清单表，逐行抽取。
    ④ 表头匹配时自动忽略星号(*)与空格：Excel 的“*课程编号”能对上 Word 的“课程编号”。
       - Word 里有的列 → 按行自动填；
       - Word 里没有的列（提交人/开课单位/课程英文名…）→ 变成统一填写框，填一次应用到所有行。
    ⑤ 一次性把 N 行追加进 Excel（表头之下）。写不进时自动另存副本。

用法：
    python gui.py
"""
import os
import re
import sys
import time
import tempfile
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import word_to_table as core
import xlsx_fill
import doc_notes
import shared_notes
import office_tools
import video_tools
import watermark_tools
import demo_page
import ppt_gen
from xlsx_fill import colkey

# 打包成无控制台 exe 后 sys.stdout/stderr 为 None，
# 任何组件往其上写一行（警告/调试）都会直接崩，这里指向空设备兜底
if sys.stdout is None:
    sys.stdout = open(os.devnull, 'w')
if sys.stderr is None:
    sys.stderr = open(os.devnull, 'w')


def _silence_libpng_stderr():
    """屏蔽第三方库(Pillow/libpng)读取带异常色彩配置文件的 PNG 时打印的
    “libpng warning: iCCP: known incorrect sRGB profile” 无害警告。
    该警告由 C 层直接写 fd 2，Python 的 warnings 拦不住，故把 fd 2 重定向到空设备；
    同时把 sys.stderr 另接到原始 stderr 的副本，保证真正的 Python 报错仍然可见。"""
    try:
        if sys.stderr is None:
            return
        saved = os.dup(2)                       # 复制真实 stderr 句柄（控制台）
        sys.stderr = os.fdopen(saved, 'w', closefd=False)
        devnull = os.open(os.devnull, os.O_WRONLY)
        os.dup2(devnull, 2)                     # C 层 stderr（含 libpng 警告）→ 空设备
        os.close(devnull)
    except Exception:                           # 无控制台 exe 等场景静默跳过
        pass


_silence_libpng_stderr()

# 个别列在 Word 里叫法不同，这里给别名（可继续补充）
ALIASES = {
    '课程性质': ['修读方式', '性质', '修读性质', '必修选修'],
    '考核方式': ['考核类型', '考试方式', '考核形式'],
    '课程类别': ['模块名称', '类别', '课程分类'],
    '学时': ['总学时', '学时数', '课内学时', '周学时', '学时合计', '理论学时'],
    '学分': ['总学分', '学分数'],
    '开课学期': ['学期', '开设学期', '所用学期', '开课学期'],
    '课程编号': ['课程代码', '课程代号', '编号', '课程序号'],
    '课程名称': ['名称', '课程中文名称'],
    '课程属性': ['属性'],
    '课程类型': ['类型'],
}

# 新生成的表格产物默认存放目录（不存在会自动创建）；
# 别的电脑没有 E 盘时，自动退到 桌面\填表工具产物，再不行退到程序所在目录

def _default_out_dir():
    if getattr(sys, 'frozen', False):          # 打包成 exe 后 = exe 所在目录
        base = os.path.dirname(sys.executable)
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    for cand in (r'E:\重要文件',
                 os.path.join(os.path.expanduser('~'), 'Desktop', '填表工具产物'),
                 os.path.join(base, '填表工具产物')):
        try:
            os.makedirs(cand, exist_ok=True)
            return cand
        except OSError:
            continue
    return os.getcwd()


OUT_DIR = _default_out_dir()


def norm(h):
    """表头归一化：去星号、去所有空白，便于 *课程编号/修 读方式 匹配。"""
    return re.sub(r'\s+', '', (h or '').replace('*', '').replace('＊', ''))


def _col_letter_to_num(s):
    n = 0
    for ch in s.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def _num_to_col_letter(n):
    s = ''
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _spec_to_num(s):
    """'E'→5；'5'→5；非字母数字 → None。"""
    s = s.strip().upper()
    if s.isdigit():
        return int(s)
    if re.fullmatch(r'[A-Z]+', s):
        return _col_letter_to_num(s)
    return None


def parse_colspec(text):
    """把用户输入的列规格解析成列字母集合。支持：
    'E' / 'E列' / '5' / 'A:C' / 'A-C' / '1-3' / 'A,B,D' / '字段标题(E列)'。
    无法解析时返回 None（调用方据此提示）。"""
    t = (text or '').strip()
    if not t:
        return None
    m = re.search(r'\(([A-Za-z]+)列?\)', t)   # 下拉标签里的括号列字母优先
    if m:
        return {m.group(1).upper()}
    t = t.replace('列', '')
    out = set()
    for tok in re.split(r'[,;、\s]+', t):
        tok = tok.strip()
        if not tok:
            continue
        rm = re.match(r'^([A-Za-z]+|\d+)\s*[:\-~～]\s*([A-Za-z]+|\d+)$', tok)
        if rm:
            a = _spec_to_num(rm.group(1))
            b = _spec_to_num(rm.group(2))
            if a and b:
                for i in range(min(a, b), max(a, b) + 1):
                    out.add(_num_to_col_letter(i))
                continue
        n = _spec_to_num(tok)
        if n:
            out.add(_num_to_col_letter(n))
    return out or None


def parse_rowspec(text):
    """行规格 → set(int)。支持 '2' / '2:10' / '2-10' / '2,5,8' / '第3行'。"""
    t = re.sub(r'\s+', '', (text or '')).replace('第', '').replace('行', '')
    if not t:
        return None
    out = set()
    for tok in re.split(r'[,;、]+', t):
        if not tok:
            continue
        rm = re.match(r'^(\d+)[:\-~～](\d+)$', tok)
        if rm:
            a, b = int(rm.group(1)), int(rm.group(2))
            out.update(range(min(a, b), max(a, b) + 1))
            continue
        if tok.isdigit():
            out.add(int(tok))
    return out or None


class App:
    def __init__(self, root):
        self.root = root
        root.title('Word 扫描 → 填写表格工具（扫课程表·多行）')
        root.geometry('960x660')

        self.header = {}        # {列字母: 表头文本}
        self.word_rows = []     # [ {源列名: 值}, ... ]
        self.col_source = {}    # {列字母: ('word', 源列名) 或 ('single', 表头文本)}
        self.single_vars = {}   # {表头文本: StringVar}
        self.map_vars = {}      # {列字母: StringVar}  列对应下拉框
        self.src_cols = []      # 表1 的列名列表
        self.tool_cols = {}     # {下拉框显示名: 列字母}
        self.var_tool_file = tk.StringVar()   # ⑤ 单独扫描的表格（空=用加工表格本身）
        self.var_tool_sheet = tk.StringVar()  # ⑤ 扫描表格选中的工作表
        self.out_dir = OUT_DIR  # 产物默认目录
        self.demo_history_path = demo_page.history_file()
        self.single_kept = {}   # {归一化列名: 用户手填值}  重建界面时不丢
        self.map_kept = {}      # {归一化目标列名: 用户选的表1列}  重扫时不丢
        self.validations = {}   # {列字母: 目标表下拉允许值}

        self._build_ui()

    # ---------------- UI ----------------
    def _build_ui(self):
        pad = {'padx': 8, 'pady': 4}

        top = ttk.LabelFrame(self.root, text='① 选择文件')
        top.pack(fill='x', **pad)
        self.var_src = tk.StringVar()
        self.var_xlsx = tk.StringVar()
        self.var_out = tk.StringVar()
        self.var_srctype = tk.StringVar(value='word')
        trow = ttk.Frame(top)
        trow.grid(row=0, column=0, columnspan=3, sticky='w', padx=4, pady=3)
        ttk.Label(trow, text='扫描源：').pack(side='left')
        ttk.Radiobutton(trow, text='Word文档', variable=self.var_srctype,
                        value='word', command=self._on_srctype).pack(side='left', padx=4)
        ttk.Radiobutton(trow, text='Excel表格(表1)', variable=self.var_srctype,
                        value='excel', command=self._on_srctype).pack(side='left', padx=4)
        self._file_row(top, '扫描的Word：', self.var_src, self.choose_src, 1)
        self._file_row(top, '表格(定扫描内容)：', self.var_xlsx, self.choose_xlsx, 2)
        self._file_row(top, '填写到(默认=上表)：', self.var_out, self.choose_out, 3)

        bar = ttk.Frame(self.root)
        bar.pack(fill='x', **pad)
        ttk.Button(bar, text='② 扫描课程', command=self.do_scan).pack(side='left', padx=4)
        ttk.Button(bar, text='③ 填写到表格', command=self.do_export).pack(side='left', padx=4)
        ttk.Button(bar, text='打开结果', command=self.open_output).pack(side='left', padx=4)
        self.var_clear = tk.BooleanVar(value=True)
        ttk.Checkbutton(bar, text='清空表头以下旧行', variable=self.var_clear).pack(
            side='left', padx=8)
        self.var_center = tk.BooleanVar(value=True)
        ttk.Checkbutton(bar, text='一键文字居中', variable=self.var_center).pack(
            side='left', padx=4)
        self.var_safe = tk.BooleanVar(value=True)
        ttk.Checkbutton(bar, text='下拉选项外的值留空', variable=self.var_safe).pack(
            side='left', padx=4)
        ttk.Label(bar, text='空值填:').pack(side='left')
        self.var_fill = tk.StringVar(value='')
        ttk.Entry(bar, textvariable=self.var_fill, width=4).pack(side='left', padx=2)
        self.lbl_status = ttk.Label(bar, text='请先选择 Word 与表格', foreground='#0a7')
        self.lbl_status.pack(side='right', padx=8)

        # 中间可滚动区域（超高用滚轮上下滑；超宽用底部横条/Shift+滚轮左右滑）
        self.canvas = tk.Canvas(self.root, highlightthickness=0)
        main_vsb = ttk.Scrollbar(self.root, orient='vertical', command=self.canvas.yview)
        main_hsb = ttk.Scrollbar(self.root, orient='horizontal', command=self.canvas.xview)
        self.canvas.configure(yscrollcommand=main_vsb.set,
                              xscrollcommand=main_hsb.set)
        main_vsb.pack(side='right', fill='y', pady=4)
        main_hsb.pack(side='bottom', fill='x', padx=8)
        self.canvas.pack(side='left', fill='both', expand=True, padx=(8, 0), pady=4)
        self.inner = ttk.Frame(self.canvas)
        self._win_id = self.canvas.create_window((0, 0), window=self.inner, anchor='nw')
        self.inner.bind('<Configure>',
                        lambda e: self.canvas.configure(scrollregion=self.canvas.bbox('all')))
        self.canvas.bind('<Configure>', self._on_canvas_resize)
        self.canvas.bind_all('<MouseWheel>', self._on_wheel)

        self.single_frame = ttk.LabelFrame(
            self.inner, text='③ 源里没有的列（统一填一次，应用到每一行；源里有的列会自动扫）')
        self.single_frame.pack(fill='x', **pad)

        self.map_frame = ttk.LabelFrame(
            self.inner, text='②·列对应设置（扫描源=Excel表格时可用：目标列 ← 表1的哪一列）')
        self.map_frame.pack(fill='x', **pad)

        mid = ttk.LabelFrame(self.inner, text='④ 扫描预览（来自 Word 内课程表）')
        mid.pack(fill='x', **pad)
        self.tree = ttk.Treeview(mid, columns=('c0',), show='headings', height=8)
        vsb = ttk.Scrollbar(mid, orient='vertical', command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side='left', fill='both', expand=True)
        vsb.pack(side='right', fill='y')

        # ---------- ⑤ 加工工具 ----------
        tools = ttk.LabelFrame(self.inner, text='⑤ 加工工具（不改变表格原格式）')
        tools.pack(fill='x', **pad)
        r0 = ttk.Frame(tools)
        r0.pack(fill='x', padx=6, pady=2)
        ttk.Label(r0, text='加工的表格：').pack(side='left')
        self.var_tool = tk.StringVar()
        ttk.Entry(r0, textvariable=self.var_tool).pack(side='left', fill='x', expand=True)
        ttk.Button(r0, text='浏览…', command=self.choose_tool).pack(side='left', padx=4)
        ttk.Button(r0, text='打开表格', command=self.open_tool).pack(side='left', padx=4)

        rt = ttk.Frame(tools)
        rt.pack(fill='x', padx=6, pady=2)
        ttk.Button(rt, text='单独扫描表格', command=self.do_scan_tool).pack(side='left')
        ttk.Label(rt, text='另一张表：').pack(side='left', padx=(8, 0))
        ttk.Entry(rt, textvariable=self.var_tool_file, width=26).pack(side='left')
        ttk.Label(rt, text='工作表：').pack(side='left', padx=(8, 0))
        self.cmb_sheet = ttk.Combobox(rt, textvariable=self.var_tool_sheet,
                                      width=14, state='readonly')
        self.cmb_sheet.pack(side='left')
        self.cmb_sheet.bind('<<ComboboxSelected>>', lambda e: self._on_sheet_change())
        self.lbl_toolinfo = ttk.Label(rt, text='（加字可改扫别的表格；不选则用上面“加工的表格”）',
                                      foreground='#666')
        self.lbl_toolinfo.pack(side='left', padx=6)

        r1 = ttk.Frame(tools)
        r1.pack(fill='x', padx=6, pady=2)
        ttk.Label(r1, text='加字范围:').pack(side='left')
        self.var_scope = tk.StringVar(value='col')
        ttk.Radiobutton(r1, text='按列', variable=self.var_scope, value='col').pack(side='left')
        self.cmb_toolcol = ttk.Combobox(r1, width=12)
        self.cmb_toolcol.pack(side='left', padx=2)
        ttk.Label(r1, text='（可输入列字母/序号/区间，如 E、A:C、1-3、A,B；也可从下拉选）',
                  foreground='#666').pack(side='left', padx=2)
        ttk.Radiobutton(r1, text='按行', variable=self.var_scope, value='row').pack(
            side='left', padx=(8, 0))
        self.var_toolrow = tk.StringVar()
        ttk.Entry(r1, textvariable=self.var_toolrow, width=10).pack(side='left', padx=2)
        ttk.Label(r1, text='（行号，可填区间如 2:10 或 2,5）',
                  foreground='#666').pack(side='left', padx=2)
        self.var_pos = tk.StringVar(value='before')
        ttk.Radiobutton(r1, text='加在前面', variable=self.var_pos, value='before').pack(
            side='left', padx=(8, 0))
        ttk.Radiobutton(r1, text='加在后面', variable=self.var_pos, value='after').pack(
            side='left')
        self.var_addtext = tk.StringVar()
        ttk.Entry(r1, textvariable=self.var_addtext, width=10).pack(side='left', padx=4)
        ttk.Button(r1, text='执行添加', command=self.do_add_text).pack(side='left', padx=4)

        r2 = ttk.Frame(tools)
        r2.pack(fill='x', padx=6, pady=2)
        ttk.Label(r2, text='切割：每份行数').pack(side='left')
        self.var_per = tk.StringVar(value='10000')
        ttk.Entry(r2, textvariable=self.var_per, width=8).pack(side='left', padx=4)
        ttk.Button(r2, text='切割表格', command=self.do_split).pack(side='left', padx=4)
        ttk.Label(r2, text='（每份保留表头与原格式，生成在同目录）').pack(side='left', padx=8)

        r3 = ttk.Frame(tools)
        r3.pack(fill='x', padx=6, pady=2)
        ttk.Label(r3, text='提取：标识包含').pack(side='left')
        self.var_mark = tk.StringVar()
        ttk.Entry(r3, textvariable=self.var_mark, width=12).pack(side='left', padx=4)
        ttk.Button(r3, text='提取含标识的行', command=self.do_extract).pack(side='left', padx=4)
        ttk.Label(r3, text='（含标识的行整行抄到新表，表头+格式不变）').pack(
            side='left', padx=8)

        # ---------- ⑥ 独立新功能 ----------
        solo = ttk.LabelFrame(self.inner, text='⑥ 新功能（独立使用，不改变原文件格式）')
        solo.pack(fill='x', **pad)

        y1 = ttk.Frame(solo)
        y1.pack(fill='x', padx=6, pady=2)
        ttk.Label(y1, text='扫描黄色标识：源表格').pack(side='left')
        self.var_ytool = tk.StringVar()
        ttk.Entry(y1, textvariable=self.var_ytool, width=20).pack(side='left', padx=4)
        ttk.Button(y1, text='浏览…', command=self.choose_ytool).pack(side='left')
        ttk.Label(y1, text='填入目标表格').pack(side='left', padx=(8, 0))
        self.var_ytarget = tk.StringVar()
        ttk.Entry(y1, textvariable=self.var_ytarget, width=20).pack(side='left', padx=4)
        ttk.Button(y1, text='浏览…', command=self.choose_ytarget).pack(side='left')
        y1b = ttk.Frame(solo)
        y1b.pack(fill='x', padx=6, pady=(0, 2))
        ttk.Button(y1b, text='扫描黄标并填入', command=self.do_scan_yellow).pack(side='left')
        ttk.Label(y1b, text='（先扫描黄底标识行，确认后按列对应填进目标表格，格式不变）').pack(
            side='left', padx=6)

        y2 = ttk.Frame(solo)
        y2.pack(fill='x', padx=6, pady=2)
        ttk.Label(y2, text='标红指定内容：文档').pack(side='left')
        self.var_reddoc = tk.StringVar()
        ttk.Entry(y2, textvariable=self.var_reddoc, width=28).pack(side='left', padx=4)
        ttk.Button(y2, text='浏览…', command=self.choose_reddoc).pack(side='left')
        y2b = ttk.Frame(solo)
        y2b.pack(fill='x', padx=6, pady=(0, 2))
        ttk.Label(y2b, text='关键字').pack(side='left')
        self.var_red = tk.StringVar()
        ttk.Entry(y2b, textvariable=self.var_red, width=12).pack(side='left', padx=4)
        self.var_inplace = tk.BooleanVar(value=False)
        ttk.Checkbutton(y2b, text='直接改原文件', variable=self.var_inplace).pack(side='left')
        ttk.Button(y2b, text='标红', command=self.do_mark_red).pack(side='left', padx=6)
        ttk.Button(y2b, text='提取标红→Word', command=self.do_extract_red).pack(side='left')
        ttk.Label(y2b, text='（把文档里的关键字标成红色字体，其它格式不变；提取则把标红所在整段/整格汇成新 Word）').pack(
            side='left', padx=4)

        # ---------- ⑦ 生成日志/笔记 ----------
        note = ttk.LabelFrame(self.inner, text='⑦ 生成日志/笔记（图片/表格/Word/文本 → 整理成文档）')
        note.pack(fill='x', **pad)

        z1 = ttk.Frame(note)
        z1.pack(fill='x', padx=6, pady=2)
        ttk.Label(z1, text='源文档：').pack(side='left')
        self.var_notes_src = tk.StringVar()
        ttk.Entry(z1, textvariable=self.var_notes_src, width=28).pack(side='left', padx=4)
        ttk.Button(z1, text='浏览…', command=self.choose_notes_src).pack(side='left')
        ttk.Label(z1, text='（可多选：图片/表格/Word/文本）').pack(side='left', padx=4)

        z2 = ttk.Frame(note)
        z2.pack(fill='x', padx=6, pady=(0, 2))
        self.var_note_kind = tk.StringVar(value='笔记')
        ttk.Label(z2, text='种类').pack(side='left')
        ttk.Radiobutton(z2, text='笔记', variable=self.var_note_kind,
                        value='笔记').pack(side='left')
        ttk.Radiobutton(z2, text='日志', variable=self.var_note_kind,
                        value='日志').pack(side='left')
        self.var_note_detail = tk.StringVar(value='简单')
        ttk.Label(z2, text='话语').pack(side='left', padx=(8, 0))
        ttk.Radiobutton(z2, text='简单', variable=self.var_note_detail,
                        value='简单').pack(side='left')
        ttk.Radiobutton(z2, text='具体', variable=self.var_note_detail,
                        value='具体').pack(side='left')
        self.var_note_fmt = tk.StringVar(value='Word')
        ttk.Label(z2, text='存为').pack(side='left', padx=(8, 0))
        ttk.Radiobutton(z2, text='Word文档', variable=self.var_note_fmt,
                        value='Word').pack(side='left')
        ttk.Radiobutton(z2, text='记事本', variable=self.var_note_fmt,
                        value='记事本').pack(side='left')
        ttk.Button(z2, text='生成', command=self.do_make_notes).pack(side='left', padx=6)
        self.var_note_who = tk.StringVar(value=shared_notes.user())
        ttk.Label(z2, text='上传人').pack(side='left', padx=(8, 0))
        ttk.Entry(z2, textvariable=self.var_note_who, width=8).pack(side='left', padx=2)
        self.var_note_tosave = tk.BooleanVar(value=True)
        ttk.Checkbutton(z2, text='生成后存入共享库',
                        variable=self.var_note_tosave).pack(side='left', padx=(8, 0))

        # 共享文档库：任何人都能上传，保存后大家打开工具都能看到
        z3 = ttk.Frame(note)
        z3.pack(fill='x', padx=6, pady=(2, 0))
        ttk.Label(z3, text='共享库：').pack(side='left')
        self.var_share_dir = tk.StringVar(value=shared_notes.lib_dir())
        ttk.Entry(z3, textvariable=self.var_share_dir, width=34).pack(side='left', padx=4)
        ttk.Button(z3, text='上传文档到共享库', command=self.do_upload_share).pack(side='left')
        ttk.Button(z3, text='换路径', command=self.choose_share_dir).pack(side='left', padx=4)
        ttk.Button(z3, text='刷新', command=self.refresh_share_lib).pack(side='left')
        ttk.Label(z3, text='（路径改成局域网共享盘，同事就能看到彼此上传）').pack(
            side='left', padx=4)

        z4 = ttk.Frame(note)
        z4.pack(fill='x', padx=6, pady=2)
        self.tree_share = ttk.Treeview(z4, columns=('f', 'w', 't', 's'),
                                       show='headings', height=5)
        for col, text, w in (('f', '文件', 260), ('w', '上传人', 90),
                             ('t', '上传时间', 120), ('s', '大小', 80)):
            self.tree_share.heading(col, text=text)
            self.tree_share.column(col, width=w, anchor='w')
        vs4 = ttk.Scrollbar(z4, orient='vertical', command=self.tree_share.yview)
        self.tree_share.configure(yscrollcommand=vs4.set)
        self.tree_share.pack(side='left', fill='both', expand=True)
        vs4.pack(side='left', fill='y')
        ops = ttk.Frame(z4)
        ops.pack(side='left', fill='y', padx=4)
        ttk.Button(ops, text='打开', command=self.open_share_item).pack(fill='x', pady=1)
        ttk.Button(ops, text='用它生成日志', command=self.notes_from_share).pack(fill='x', pady=1)
        ttk.Button(ops, text='删除', command=self.delete_share_item).pack(fill='x', pady=1)
        ttk.Button(ops, text='查看占用', command=self.show_share_usage).pack(fill='x', pady=1)
        self.share_items = []
        self.refresh_share_lib()

        # 文字笔记：直接写一段存进共享库，大家都能看到
        z5 = ttk.Frame(note)
        z5.pack(fill='x', padx=6, pady=(2, 4))
        ttk.Label(z5, text='文字笔记（写完存入共享库，大家可见）：').pack(
            side='left', anchor='n', padx=(0, 4))
        self.txt_note = tk.Text(z5, height=3, width=40, font=('Microsoft YaHei', 10))
        self.txt_note.pack(side='left', fill='x', expand=True)
        nb = ttk.Frame(z5)
        nb.pack(side='left', fill='y', padx=4)
        ttk.Button(nb, text='存入共享库', command=self.do_save_text_note).pack(
            fill='x', pady=1)

        # ---------- ⑧ 投标演示页面生成 ----------
        demo = ttk.LabelFrame(self.inner, text='⑧ 投标演示页面生成（同一内容最多 3 套不同方案）')
        demo.pack(fill='x', **pad)

        d1 = ttk.Frame(demo)
        d1.pack(fill='x', padx=6, pady=2)
        ttk.Label(d1, text='参数文件：').pack(side='left')
        self.var_demo_file = tk.StringVar()
        ttk.Entry(d1, textvariable=self.var_demo_file, width=28).pack(side='left', padx=4)
        ttk.Button(d1, text='浏览…', command=self.choose_demo_file).pack(side='left')
        ttk.Label(d1, text='（Word/表格/文本，载入后内容进下方框，可直接改）').pack(
            side='left', padx=4)

        d2 = ttk.Frame(demo)
        d2.pack(fill='x', padx=6, pady=2)
        self.txt_demo = tk.Text(d2, height=7, font=('Microsoft YaHei', 10))
        self.txt_demo.pack(fill='x')

        d3 = ttk.Frame(demo)
        d3.pack(fill='x', padx=6, pady=(0, 2))
        ttk.Button(d3, text='生成演示页面', command=self.do_make_demo).pack(side='left')
        ttk.Label(d3, text='（前三次依次生成晴空蓝 / 星云紫 / 青岚绿；第四次起循环复用三套风格继续生成，重开工具仍记次）').pack(
            side='left', padx=6)
        ttk.Label(demo, text='导航与业务环节随参数内容变化；每套独立保存为 HTML，双击打开。').pack(
            anchor='w', padx=8, pady=(0, 5))

        writer = ttk.LabelFrame(self.inner, text='⑨ 项目文书生成（上传资料＋提示词 → Word / Excel）')
        writer.pack(fill='x', **pad)
        ttk.Button(writer, text='打开项目文书工作台', command=self.open_project_writer).pack(
            side='left', padx=8, pady=8)
        ttk.Label(writer, text='申报书 / 论证报告 / 实施方案 / 绩效目标表；可选样式参考和联网查证。\n'
                  '需配置自己的 AI 接口；资料只读，结果另存，不受演示页面三次限制。').pack(
            side='left', padx=8, pady=6)

        # ---------- ⑩ 生成PPT演示文稿 ----------
        ppt = ttk.LabelFrame(self.inner, text='⑩ 生成PPT（Word文档或文字描述 → 演示文稿）')
        ppt.pack(fill='x', **pad)

        p1 = ttk.Frame(ppt)
        p1.pack(fill='x', padx=6, pady=2)
        ttk.Label(p1, text='来源文件：').pack(side='left')
        self.var_ppt_file = tk.StringVar()
        ttk.Entry(p1, textvariable=self.var_ppt_file, width=28).pack(side='left', padx=4)
        ttk.Button(p1, text='浏览…', command=self.choose_ppt_file).pack(side='left')
        ttk.Label(p1, text='（Word/表格/文本，载入后内容进下方框，可直接改）').pack(
            side='left', padx=4)

        p2 = ttk.Frame(ppt)
        p2.pack(fill='x', padx=6, pady=2)
        self.txt_ppt = tk.Text(p2, height=6, font=('Microsoft YaHei', 10))
        self.txt_ppt.pack(fill='x')

        p3 = ttk.Frame(ppt)
        p3.pack(fill='x', padx=6, pady=(0, 2))
        ttk.Label(p3, text='标题').pack(side='left')
        self.var_ppt_title = tk.StringVar()
        ttk.Entry(p3, textvariable=self.var_ppt_title, width=20).pack(side='left', padx=4)
        ttk.Label(p3, text='配色').pack(side='left', padx=(8, 0))
        self.var_ppt_scheme = tk.StringVar(value='自动轮流')
        ttk.Combobox(p3, textvariable=self.var_ppt_scheme, state='readonly', width=8,
                     values=('自动轮流', '晴空蓝', '星云紫', '青岚绿')).pack(side='left', padx=4)
        ttk.Label(p3, text='页数').pack(side='left', padx=(8, 0))
        self.var_ppt_pages = tk.StringVar(value='10')
        ttk.Spinbox(p3, textvariable=self.var_ppt_pages, from_=1, to=30, width=4).pack(
            side='left', padx=4)
        ttk.Label(p3, text='页（1~30）').pack(side='left')
        ttk.Button(p3, text='生成PPT', command=self.do_make_ppt).pack(side='left', padx=6)
        ttk.Label(ppt, text='（页标题行用##开头、要点行用-开头；也可直接粘贴一段需求描述（含“主题是/面向/几页”等），'
                  '会自动抽主题与章节、每页扩成6~7条要点+小结+演讲备注并自动绘制矢量配图（要点自动铺满整页）；页数不够自动拆页/补充，多了自动合并）').pack(
            anchor='w', padx=8, pady=(0, 5))

        # ---------- ⑪ 多表合并汇总 ----------
        mrg = ttk.LabelFrame(self.inner, text='⑪ 多表合并汇总（多份同格式 Excel → 一张总表）')
        mrg.pack(fill='x', **pad)
        self.var_merge_files = tk.StringVar()
        self._frow(mrg, '要合并的表格：', self.var_merge_files, self.choose_merge_files,
                   '（可多选，表头需一致）')
        fr = ttk.Frame(mrg); fr.pack(fill='x', padx=6, pady=(0, 4))
        self.var_merge_dedup = tk.BooleanVar(value=False)
        ttk.Checkbutton(fr, text='全行相同只留一条（去重）',
                        variable=self.var_merge_dedup).pack(side='left')
        ttk.Button(fr, text='合并汇总', command=self.do_merge).pack(side='left', padx=8)

        # ---------- ⑫ 批量套打 / 邮件合并 ----------
        mm = ttk.LabelFrame(self.inner, text='⑫ 批量套打/邮件合并（Word模板 {{列名}} + 名单表 → N 份）')
        mm.pack(fill='x', **pad)
        self.var_mm_tpl = tk.StringVar(); self.var_mm_data = tk.StringVar(); self.var_mm_name = tk.StringVar()
        self._frow(mm, '模板文档：', self.var_mm_tpl, self.choose_mm_tpl,
                   '（.docx/.doc，内含 {{姓名}} 之类占位符）')
        self._frow(mm, '名单表格：', self.var_mm_data, self.choose_mm_data,
                   '（.xlsx，表头=占位符名）')
        fr = ttk.Frame(mm); fr.pack(fill='x', padx=6, pady=(0, 4))
        ttk.Label(fr, text='用哪一列给文件命名：').pack(side='left')
        ttk.Entry(fr, textvariable=self.var_mm_name, width=12).pack(side='left', padx=4)
        ttk.Button(fr, text='批量生成', command=self.do_mailmerge).pack(side='left', padx=8)

        # ---------- ⑬ 两表对账 ----------
        rec = ttk.LabelFrame(self.inner, text='⑬ 两表对账（按关键列找 只在A/只在B/值不同）')
        rec.pack(fill='x', **pad)
        self.var_rec_a = tk.StringVar(); self.var_rec_b = tk.StringVar(); self.var_rec_key = tk.StringVar()
        self._frow(rec, '表A：', self.var_rec_a, self.choose_rec_a)
        self._frow(rec, '表B：', self.var_rec_b, self.choose_rec_b)
        fr = ttk.Frame(rec); fr.pack(fill='x', padx=6, pady=(0, 4))
        ttk.Label(fr, text='关键列名：').pack(side='left')
        ttk.Entry(fr, textvariable=self.var_rec_key, width=14).pack(side='left', padx=4)
        ttk.Button(fr, text='开始对账', command=self.do_reconcile).pack(side='left', padx=8)

        # ---------- ⑭ 格式清洗 ----------
        cln = ttk.LabelFrame(self.inner, text='⑭ 格式清洗（Excel 去空格/全角转半角，不动数值与格式）')
        cln.pack(fill='x', **pad)
        self.var_clean_file = tk.StringVar()
        self._frow(cln, '要清洗的表格：', self.var_clean_file, self.choose_clean_file)
        fr = ttk.Frame(cln); fr.pack(fill='x', padx=6, pady=(0, 4))
        ttk.Button(fr, text='清洗并另存', command=self.do_clean).pack(side='left', padx=4)

        # ---------- ⑮ 文档互转 + 批量重命名 ----------
        cvt = ttk.LabelFrame(self.inner, text='⑮ 文档互转 + 批量重命名')
        cvt.pack(fill='x', **pad)
        self.var_conv_files = tk.StringVar()
        self._frow(cvt, '要转换的文件：', self.var_conv_files, self.choose_conv_files,
                   '（.doc/.wps→docx，.xlsx→csv）')
        fr = ttk.Frame(cvt); fr.pack(fill='x', padx=6, pady=2)
        ttk.Button(fr, text='转换到产物目录', command=self.do_convert).pack(side='left', padx=4)
        self.var_ren_dir = tk.StringVar()
        self._frow(cvt, '重命名文件夹：', self.var_ren_dir, self.choose_ren_dir)
        fr = ttk.Frame(cvt); fr.pack(fill='x', padx=6, pady=(0, 4))
        self.var_ren_mode = tk.StringVar(value='replace')
        ttk.Radiobutton(fr, text='查找替换', variable=self.var_ren_mode,
                        value='replace').pack(side='left')
        ttk.Radiobutton(fr, text='加前缀+序号', variable=self.var_ren_mode,
                        value='prefix').pack(side='left', padx=6)
        self.var_ren_find = tk.StringVar(); self.var_ren_repl = tk.StringVar()
        self.var_ren_prefix = tk.StringVar(value='文件_'); self.var_ren_ext = tk.StringVar(value='.docx')
        ttk.Label(fr, text='查找').pack(side='left', padx=(6, 0))
        ttk.Entry(fr, textvariable=self.var_ren_find, width=8).pack(side='left', padx=2)
        ttk.Label(fr, text='替换').pack(side='left')
        ttk.Entry(fr, textvariable=self.var_ren_repl, width=8).pack(side='left', padx=2)
        ttk.Label(fr, text='前缀').pack(side='left')
        ttk.Entry(fr, textvariable=self.var_ren_prefix, width=8).pack(side='left', padx=2)
        ttk.Label(fr, text='类型').pack(side='left')
        ttk.Entry(fr, textvariable=self.var_ren_ext, width=6).pack(side='left', padx=2)
        ttk.Button(fr, text='预览并重命名', command=self.do_rename).pack(side='left', padx=8)

        # ---------- ⑯ 周报 / 会议纪要 ----------
        wk = ttk.LabelFrame(self.inner, text='⑯ 周报 / 会议纪要（贴要点 → 规范 Word → 存共享库）')
        wk.pack(fill='x', **pad)
        fr = ttk.Frame(wk); fr.pack(fill='x', padx=6, pady=2)
        self.var_wk_kind = tk.StringVar(value='周报')
        ttk.Radiobutton(fr, text='周报', variable=self.var_wk_kind,
                        value='周报').pack(side='left')
        ttk.Radiobutton(fr, text='会议纪要', variable=self.var_wk_kind,
                        value='会议纪要').pack(side='left', padx=6)
        self.var_wk_who = tk.StringVar(value=shared_notes.user())
        ttk.Label(fr, text='记录人').pack(side='left', padx=(8, 0))
        ttk.Entry(fr, textvariable=self.var_wk_who, width=8).pack(side='left', padx=2)
        ttk.Button(fr, text='生成并存共享库', command=self.do_weekly).pack(side='left', padx=8)
        self.txt_wk = tk.Text(wk, height=4, font=('Microsoft YaHei', 10))
        self.txt_wk.pack(fill='x', padx=6, pady=(0, 2))
        ttk.Label(wk, text='（每行一条要点，自动分“本周完成/下周计划”或“议题/结论与待办”）').pack(
            anchor='w', padx=8, pady=(0, 4))

        # ---------- ⬓ 视频合成 ----------
        self.video_files = []
        vd = ttk.LabelFrame(self.inner, text='⬓ 视频合成（选 2-%d 个视频，按顺序拼成一个 .mp4）'
                            % video_tools.MAX_VIDEOS)
        vd.pack(fill='x', **pad)
        bfr = ttk.Frame(vd); bfr.pack(fill='x', padx=6, pady=(2, 0))
        ttk.Button(bfr, text='添加视频', command=self.add_videos).pack(side='left', padx=2)
        ttk.Button(bfr, text='移除选中', command=self.remove_video).pack(side='left', padx=2)
        ttk.Button(bfr, text='上移', command=lambda: self.move_video(-1)).pack(side='left', padx=2)
        ttk.Button(bfr, text='下移', command=lambda: self.move_video(1)).pack(side='left', padx=2)
        ttk.Button(bfr, text='清空', command=self.clear_videos).pack(side='left', padx=2)
        self.lst_video = tk.Listbox(vd, height=5, exportselection=False,
                                    font=('Microsoft YaHei', 9))
        self.lst_video.pack(fill='x', padx=6, pady=2)
        opt = ttk.Frame(vd); opt.pack(fill='x', padx=6, pady=(0, 2))
        ttk.Label(opt, text='输出分辨率').pack(side='left')
        self.var_video_res = tk.StringVar(value='保持原始')
        ttk.Combobox(opt, textvariable=self.var_video_res, width=12, state='readonly',
                     values=['保持原始', '放大到1080P', '放大到2K', '放大到4K(慢)']).pack(side='left', padx=4)
        ttk.Label(opt, text='（画质用下面两个按钮选）').pack(side='left', padx=6)
        fr = ttk.Frame(vd); fr.pack(fill='x', padx=6, pady=(0, 4))
        ttk.Label(fr, text='输出文件名').pack(side='left')
        self.var_video_out = tk.StringVar(value='合成视频.mp4')
        ttk.Entry(fr, textvariable=self.var_video_out, width=22).pack(side='left', padx=4)
        self.btn_video = ttk.Button(fr, text='输出原画质',
                                    command=lambda: self.do_video_merge('original'))
        self.btn_video.pack(side='left', padx=6)
        self.btn_video_hd = ttk.Button(fr, text='输出高清画质',
                                        command=lambda: self.do_video_merge('hd'))
        self.btn_video_hd.pack(side='left', padx=4)
        self.lbl_video = ttk.Label(fr, text='', foreground='#08979c')
        self.lbl_video.pack(side='left', padx=6)
        ttk.Label(vd, text='（输出原画质：同格式无损直拼最快最清晰；输出高清画质：强制高画质重编码。分辨率可另选放大到1080P/2K/4K）').pack(
            anchor='w', padx=8, pady=(0, 4))

        # ---------- ⬔ 去水印 ----------
        self.wm_path = None
        self.wm_box = None
        wm = ttk.LabelFrame(self.inner, text='⬔ 去水印（照片/视频 · 框选水印位置 · 本地或网址）')
        wm.pack(fill='x', **pad)
        r1 = ttk.Frame(wm); r1.pack(fill='x', padx=6, pady=2)
        self.var_wm_src = tk.StringVar()
        ttk.Label(r1, text='文件/网址：').pack(side='left')
        ttk.Entry(r1, textvariable=self.var_wm_src, width=46).pack(side='left', padx=4)
        ttk.Button(r1, text='选文件', command=self.choose_wm_file).pack(side='left', padx=2)
        ttk.Button(r1, text='用网址', command=self.wm_use_url).pack(side='left', padx=2)
        self.lbl_wm_src = ttk.Label(r1, text='未选择', foreground='#888')
        self.lbl_wm_src.pack(side='left', padx=6)
        r2 = ttk.Frame(wm); r2.pack(fill='x', padx=6, pady=2)
        ttk.Button(r2, text='① 框选水印位置', command=self.wm_pick_box).pack(side='left', padx=2)
        self.lbl_wm_box = ttk.Label(r2, text='未框选区域', foreground='#888')
        self.lbl_wm_box.pack(side='left', padx=6)
        r3 = ttk.Frame(wm); r3.pack(fill='x', padx=6, pady=(0, 4))
        ttk.Label(r3, text='方式').pack(side='left')
        self.var_wm_method = tk.StringVar(value='智能修补')
        ttk.Combobox(r3, textvariable=self.var_wm_method, width=9, state='readonly',
                     values=['智能修补', '淡化去除']).pack(side='left', padx=4)
        ttk.Label(r3, text='输出名').pack(side='left', padx=(8, 0))
        self.var_wm_out = tk.StringVar(value='去水印结果')
        ttk.Entry(r3, textvariable=self.var_wm_out, width=16).pack(side='left', padx=4)
        self.btn_wm = ttk.Button(r3, text='② 去除水印', command=self.do_watermark)
        self.btn_wm.pack(side='left', padx=8)
        self.lbl_wm = ttk.Label(r3, text='', foreground='#08979c')
        self.lbl_wm.pack(side='left', padx=6)
        ttk.Label(wm, text='（照片用智能修补效果好；视频用淡化去除，适合固定位置台标/角标。满屏半透明水印无法干净去除）').pack(
            anchor='w', padx=8, pady=(0, 4))


    def open_project_writer(self):
        from project_ui import ProjectWindow
        window = getattr(self, 'project_writer', None)
        if window is not None and not window.closed:
            window.win.deiconify()
            window.win.lift()
            return
        self.project_writer = ProjectWindow(self.root, self.out_dir)

    def _on_canvas_resize(self, e):
        """内层宽度 = 可视宽与内容需求宽的较大者，超宽时才能横向滚到。"""
        w = max(e.width, self.inner.winfo_reqwidth())
        self.canvas.itemconfigure(self._win_id, width=w)
        self.canvas.configure(scrollregion=self.canvas.bbox('all'))

    def _on_wheel(self, e):
        """滚轮：预览树内滚树自己；按住 Shift 横向滚；其余位置纵向滚主区域。"""
        w = e.widget
        if isinstance(w, str):        # 控件已销毁时 e.widget 退化成路径字符串
            try:
                w = self.root.nametowidget(w)
            except (KeyError, tk.TclError):
                return
        if not hasattr(w, 'winfo_toplevel') or w.winfo_toplevel() is not self.root:
            return
        step = -1 * (e.delta // 120)
        if w is self.tree or str(w).startswith(str(self.tree)):
            self.tree.yview_scroll(step, 'units')
        elif getattr(e, 'state', 0) & 0x0001:   # Shift+滚轮 = 左右滑
            self.canvas.xview_scroll(step, 'units')
        else:
            self.canvas.yview_scroll(step, 'units')

    def _file_row(self, parent, label, var, cmd, row):
        lab = ttk.Label(parent, text=label, width=18)
        lab.grid(row=row, column=0, sticky='e', padx=4, pady=3)
        if row == 1:
            self.lbl_src = lab
        ttk.Entry(parent, textvariable=var).grid(row=row, column=1, sticky='we', padx=4, pady=3)
        ttk.Button(parent, text='浏览…', command=cmd).grid(row=row, column=2, padx=4, pady=3)
        parent.columnconfigure(1, weight=1)

    # ---------------- 文件选择 ----------------
    def _on_srctype(self):
        self.lbl_src.config(text='扫描的表格(表1)：' if self.var_srctype.get() == 'excel'
                            else '扫描的Word：')
        self.set_status('扫描源已切换，请重新选择文件并扫描')

    def _initial_dir(self):
        """文件对话框默认打开到桌面（兼容 OneDrive 同步的桌面），其次才回退到用户目录。"""
        home = os.path.expanduser('~')
        for d in (os.path.join(home, 'OneDrive', 'Desktop'),
                  os.path.join(home, 'Desktop'),
                  os.path.join(home, '桌面'),
                  home):
            if os.path.isdir(d):
                return d
        return home

    def choose_src(self):
        if self.var_srctype.get() == 'excel':
            p = filedialog.askopenfilename(title='选择要扫描的表格(表1)',
                                           initialdir=self._initial_dir(),
                                           filetypes=[('Excel 表格', '*.xlsx'), ('所有文件', '*.*')])
        else:
            p = filedialog.askopenfilename(title='选择要扫描的 Word 文档',
                                           initialdir=self._initial_dir(),
                                           filetypes=[('Word 文档', '*.docx *.doc *.wps'),
                                                      ('所有文件', '*.*')])
        if p:
            self.var_src.set(p)

    def choose_xlsx(self):
        p = filedialog.askopenfilename(title='选择表格（表头决定扫描什么内容）',
                                       initialdir=self._initial_dir(),
                                       filetypes=[('Excel 表格', '*.xlsx'), ('所有文件', '*.*')])
        if p:
            self.var_xlsx.set(p)
            if not self.var_out.get():
                self.var_out.set(os.path.join(self.out_dir, os.path.basename(p)))
            if not self.var_tool.get():
                self._apply_tool_file(p)
            self.set_status('已选表格，点“② 扫描课程”')

    def _apply_tool_file(self, p):
        """设置“加工的表格”并同步“工作表”下拉（默认第一张有数据的表）+刷新列。
        无论用户是点⑤的浏览还是从主表格自动带过来，都走这里，避免下拉为空。"""
        self.var_tool.set(p)
        try:
            names, default = xlsx_fill.sheet_info(p)
            self.cmb_sheet['values'] = names
            self.var_tool_sheet.set(default if default in names else
                                    (names[0] if names else ''))
        except Exception:
            self.cmb_sheet['values'] = []
            self.var_tool_sheet.set('')
        self._refresh_tool_cols()

    def choose_tool(self):
        p = filedialog.askopenfilename(title='选择要加工的表格',
                                       initialdir=self._initial_dir(),
                                       filetypes=[('Excel 表格', '*.xlsx'), ('所有文件', '*.*')])
        if p:
            self._apply_tool_file(p)

    def choose_out(self):
        p = filedialog.asksaveasfilename(title='填写到哪个表格', defaultextension='.xlsx',
                                         initialdir=self._initial_dir(),
                                         filetypes=[('Excel 表格', '*.xlsx')])
        if p:
            self.var_out.set(p)

    # ---------------- 扫描 ----------------
    def do_scan(self):
        src = self.var_src.get().strip()
        xlsx = self.var_xlsx.get().strip()
        if not src or not os.path.exists(src):
            messagebox.showwarning('提示', '请先选择要扫描的文件。')
            return
        # 选的是 Excel 表格却还停在 Word 模式 → 自动切到“Excel 表格(表1)”扫描源
        if self.var_srctype.get() == 'word' and src.lower().endswith(('.xlsx', '.xlsm')):
            self.var_srctype.set('excel')
            self._on_srctype()
            self.set_status('检测到表格文件，已自动切换为 Excel 扫描源')
        if self.var_srctype.get() != 'excel' and \
                os.path.splitext(src)[1].lower() not in ('.docx', '.doc', '.wps'):
            messagebox.showwarning(
                '格式不支持',
                '请选择 Word 文档（.docx / .doc / .wps）。\n'
                '若要用 Excel 表格作为扫描源，请选上方“Excel 表格(表1)”。')
            return
        if not xlsx or not os.path.exists(xlsx):
            messagebox.showwarning('提示', '请先选择表格。')
            return
        try:
            self.header = xlsx_fill.read_header(xlsx)
            try:
                self.validations = xlsx_fill.read_validations(xlsx)
            except Exception:
                self.validations = {}
            if self.var_srctype.get() == 'excel':
                h1, r1 = xlsx_fill.read_rows(src)
                self.src_cols = [h1[c] for c in sorted(h1, key=colkey)]
                self.word_rows = [{h1[c]: row[c] for c in row if c in h1} for row in r1]
                word_cols = set(self.src_cols)
                if not self.word_rows:
                    messagebox.showwarning('提示', '表1 里没读到数据行。')
                    return
            else:
                self.src_cols = []
                word_cols, self.word_rows = core.scan_courses(src)
        except Exception as e:
            messagebox.showerror('扫描失败', str(e))
            return
        if not self.header:
            messagebox.showwarning('提示', '表格没读到表头。')
            return
        if not self.word_rows:
            messagebox.showwarning('提示', 'Word 里没找到含“课程编号/课程名称”的课程表。')
            return

        self._map_columns(src, word_cols)
        self._build_mapping()
        self._build_single_inputs(src)
        self._refresh_tree()
        n_auto = sum(1 for v in self.col_source.values() if v[0] == 'word')
        n_single = len(self.single_vars)
        msg = '扫描到 %d 行；自动列 %d 个、需手填列 %d 个' % (
            len(self.word_rows), n_auto, n_single)
        warn = False
        bad = {}
        for cl, allowed in self.validations.items():
            if cl not in self.col_source:
                continue
            vals = set()
            for row in self.word_rows:
                v = self._cell_raw(row, cl)
                if v and v not in allowed:
                    vals.add(v)
            if vals:
                bad[norm(self.header[cl])] = vals
        if bad:
            warn = True
            tail = '，已留空' if self.var_safe.get() else '，仍照源填写'
            msg += '；注意：' + '；'.join(
                '%s 源里的 %s 不在表格下拉选项内%s' % (
                    k, '/'.join(sorted(v)[:3]), tail) for k, v in bad.items())
        self.set_status(msg, warn=warn)

    # ---------------- 列对应设置（表1 → 表2） ----------------
    def _build_mapping(self):
        """扫描源=Excel 时，为每个目标列生成下拉框选择表1的对应列。"""
        for w in self.map_frame.winfo_children():
            w.destroy()
        self.map_vars = {}
        if self.var_srctype.get() != 'excel' or not self.src_cols:
            ttk.Label(self.map_frame, text='（扫描源选“Excel表格(表1)”后，这里可设置列对应关系）'
                      ).grid(row=0, column=0, padx=6, pady=3)
            return
        opts = ['（不填）'] + self.src_cols
        col = 0
        rowi = 0
        for colletter in sorted(self.header, key=colkey):
            kind, payload = self.col_source[colletter]
            kept = self.map_kept.get(norm(self.header[colletter]))
            if kept in opts:
                init = kept                      # 用户之前手动选过的，重扫不丢
            else:
                init = payload if kind == 'word' else '（不填）'
            var = tk.StringVar(value=init)
            var.trace_add('write', lambda *a, c=colletter: self._on_map_change(c))
            ttk.Label(self.map_frame, text=norm(self.header[colletter]) + '←').grid(
                row=rowi, column=col, padx=4, pady=2, sticky='e')
            ttk.Combobox(self.map_frame, textvariable=var, values=opts,
                         width=12, state='readonly').grid(row=rowi, column=col + 1, padx=4, pady=2)
            self.map_vars[colletter] = var
            col += 2
            if col >= 6:
                col = 0
                rowi += 1

    def _on_map_change(self, colletter):
        val = self.map_vars[colletter].get()
        self.map_kept[norm(self.header[colletter])] = val   # 记住用户的选择
        if val == '（不填）':
            self.col_source[colletter] = ('single', self.header[colletter])
        else:
            self.col_source[colletter] = ('word', val)
        self._build_single_inputs(self.var_src.get().strip())
        self._refresh_tree()

    def _map_columns(self, src, word_cols):
        """把 Excel 每一列映射到 Word 列（归一化+别名）或标记为手填。"""
        word_norm = {norm(k): k for k in word_cols}
        self.col_source = {}
        for col, text in self.header.items():
            n = norm(text)
            if not n:
                self.col_source[col] = ('blank', col)   # 表头只有星号等无名列
                continue
            if n in word_norm:
                self.col_source[col] = ('word', word_norm[n])
                continue
            hit = None
            for alias in ALIASES.get(n, []):
                if alias in word_norm:
                    hit = word_norm[alias]
                    break
            if hit:
                self.col_source[col] = ('word', hit)
            else:
                self.col_source[col] = ('single', text)

    def _build_single_inputs(self, src):
        for k, v in self.single_vars.items():     # 先存下用户已填的值
            self.single_kept[norm(k)] = v.get()
        for w in self.single_frame.winfo_children():
            w.destroy()
        self.single_vars = {}
        stem = os.path.splitext(os.path.basename(src))[0]
        col = 0
        rowi = 0
        seen = set()
        for colletter in sorted(self.col_source, key=colkey):
            kind, payload = self.col_source[colletter]
            if kind != 'single' or payload in seen:
                continue
            seen.add(payload)
            n = norm(payload)
            if n in self.single_kept:             # 用用户填过的（包括清空）
                default = self.single_kept[n]
            else:
                default = stem if n == '提交人' else ''
            var = tk.StringVar(value=default)
            ttk.Label(self.single_frame, text=norm(payload) + '：').grid(
                row=rowi, column=col, padx=4, pady=3, sticky='e')
            ttk.Entry(self.single_frame, textvariable=var, width=18).grid(
                row=rowi, column=col + 1, padx=4, pady=3)
            self.single_vars[payload] = var
            col += 2
            if col >= 6:          # 每行放 3 组，超出换行
                col = 0
                rowi += 1

    def _cell_raw(self, row, colletter):
        kind, payload = self.col_source[colletter]
        if kind == 'word':
            return row.get(payload, '')
        if kind == 'blank':
            return ''
        v = self.single_vars.get(payload)
        return v.get() if v else ''

    def _cell(self, row, colletter):
        val = self._cell_raw(row, colletter)
        allowed = self.validations.get(colletter)   # 目标表有下拉限制的列
        if self.var_safe.get() and allowed and val and val not in allowed:
            return ''                               # 脏值（如代码29）不写进去
        return val

    def _refresh_tree(self):
        self.tree.delete(*self.tree.get_children())
        letters = sorted(self.header, key=colkey)
        self.tree['columns'] = tuple('c%d' % i for i in range(len(letters)))
        for i, cl in enumerate(letters):
            self.tree.heading('c%d' % i, text=norm(self.header[cl]))
            self.tree.column('c%d' % i, width=140, anchor='w')
        for row in self.word_rows:
            self.tree.insert('', 'end', values=[self._cell(row, cl) for cl in letters])

    # ---------------- 填写 ----------------
    def do_export(self):
        xlsx = self.var_xlsx.get().strip()
        out = self.var_out.get().strip() or xlsx
        if not self.word_rows:
            messagebox.showwarning('提示', '请先点“② 扫描课程”。')
            return
        rows = []
        for row in self.word_rows:
            rv = {cl: self._cell(row, cl) for cl in self.header}
            rows.append(rv)
        clear = bool(self.var_clear.get())
        center = bool(self.var_center.get())
        fill = self.var_fill.get().strip() or None
        try:
            d = os.path.dirname(os.path.abspath(out))
            if d and not os.path.exists(d):
                os.makedirs(d, exist_ok=True)
        except OSError:
            pass
        try:
            start, end = xlsx_fill.append_rows(xlsx, out, rows, clear=clear,
                                               center=center, fill=fill)
            tail = ''
        except (PermissionError, OSError) as e:
            base, ext = os.path.splitext(out)
            fallback = base + '-已填写' + ext
            try:
                start, end = xlsx_fill.append_rows(xlsx, fallback, rows, clear=clear,
                                                   center=center, fill=fill)
            except Exception as e2:
                messagebox.showerror('填写失败',
                                     '写不进目标表格：%s\n另存副本也失败：%s\n'
                                     '请先关闭正在 WPS/Excel 中打开的表格，再重试。' % (e, e2))
                return
            out = fallback
            tail = '\n注意：原表格正被打开或无权限，已改存到副本。'
        except Exception as e:
            messagebox.showerror('填写失败', str(e))
            return
        same = os.path.normpath(out) == os.path.normpath(xlsx)
        self.var_out.set(out)   # 让“打开结果”指向真正写入的文件（含 -已填写 副本）
        self.set_status('已填写第 %d~%d 行（共%d门）' % (start, end, len(rows)))
        messagebox.showinfo('完成', '已填写第 %d ~ %d 行，共 %d 门课程。\n%s%s%s'
                            % (start, end, len(rows), out,
                               '\n(即所选表格本身)' if same else '', tail))

    def open_output(self):
        out = self.var_out.get().strip()
        if out and os.path.exists(out):
            os.startfile(out)
        else:
            messagebox.showwarning('提示', '结果文件还不存在。')

    # ---------------- ⑤ 加工工具 ----------------
    def _tool_target(self):
        """加字/切割/提取的实际目标：优先“单独扫描表格”，否则“加工的表格”。
        工作表以下拉为准；下拉为空（如手输路径）时自动选第一张有数据的表，
        ⑤总是改有内容的表，而非空的模板骨架表。"""
        f = self.var_tool_file.get().strip()
        if not (f and os.path.exists(f)):
            f = self.var_tool.get().strip()
        sheet = self.var_tool_sheet.get().strip()
        if not sheet and f and os.path.exists(f):
            try:
                _, sheet = xlsx_fill.sheet_info(f)
            except Exception:
                sheet = ''
        return f, (sheet or None)

    def do_scan_tool(self):
        """选另一张表格来扫描与加字（不影响上方“加工的表格”）。"""
        p = filedialog.askopenfilename(title='选择要单独扫描的表格',
                                       initialdir=self._initial_dir(),
                                       filetypes=[('Excel 表格', '*.xlsx'), ('所有文件', '*.*')])
        if not p:
            return
        try:
            names, default = xlsx_fill.sheet_info(p)
        except Exception as e:
            messagebox.showerror('读表失败', str(e))
            return
        self.var_tool_file.set(p)
        self.cmb_sheet['values'] = names
        self.var_tool_sheet.set(default if default in names else
                                (names[0] if names else ''))
        self._refresh_tool_cols()
        self.lbl_toolinfo.config(text='已扫描：%s（工作表：%s）' % (
            os.path.basename(p), self.var_tool_sheet.get() or '—'))
        self.set_status('已单独扫描表格，共 %d 列；选列、填字后点“执行添加”' % len(self.tool_cols))

    def _on_sheet_change(self):
        self._refresh_tool_cols()

    def _refresh_tool_cols(self):
        tool, sheet = self._tool_target()
        self.tool_cols = {}
        if not tool or not os.path.exists(tool):
            self.cmb_toolcol['values'] = []
            return
        try:
            header = xlsx_fill.read_header(tool, sheet)
        except Exception as e:
            messagebox.showerror('读表失败', str(e))
            return
        opts = []
        for c in sorted(header, key=colkey):
            label = '%s(%s列)' % (norm(header[c]) or '无名', c)
            self.tool_cols[label] = c
            opts.append(label)
        self.cmb_toolcol['values'] = opts
        if opts:
            self.cmb_toolcol.current(0)

    def do_add_text(self):
        tool, sheet = self._tool_target()
        text = self.var_addtext.get()
        if not tool or not os.path.exists(tool):
            messagebox.showwarning('提示', '请先选择要加工的表格（或点“单独扫描表格”）。')
            return
        if not text:
            messagebox.showwarning('提示', '请先输入要添加的文字。')
            return
        col = row = None
        if self.var_scope.get() == 'col':
            raw = self.cmb_toolcol.get().strip()
            if not raw:
                messagebox.showwarning('提示', '请输入或选择要加字的列。')
                return
            if raw in self.tool_cols:                 # 直接选了下拉里的列标签
                col = {self.tool_cols[raw]}
            else:                                     # 手输：列字母/序号/区间
                col = parse_colspec(raw)
            if not col:
                messagebox.showwarning(
                    '无法识别列', '“%s”不是有效列。\n可输入列字母(E)、序号(5)、区间(A:C 或 1-3)或列表(A,B)。'
                    % raw)
                return
        else:
            raw = self.var_toolrow.get().strip()
            row = parse_rowspec(raw)
            if not row:
                messagebox.showwarning(
                    '无法识别行', '“%s”不是有效行号。\n可输入单行(2)、区间(2:10 或 2-10)或列表(2,5)。'
                    % raw)
                return
        before = self.var_pos.get() == 'before'
        try:
            n = xlsx_fill.add_text(tool, tool, text, before=before, col=col, row=row,
                                   sheet=sheet)
        except (PermissionError, OSError):
            messagebox.showerror('失败', '表格正被 WPS/Excel 打开，请先关闭再试。')
            return
        except Exception as e:
            messagebox.showerror('失败', str(e))
            return
        if n == 0:
            self.set_status('该列/行没有可改的文字（0 个）：请确认选对了工作表与列', warn=True)
            messagebox.showwarning(
                '修改 0 个单元格',
                '工作表“%s”的所选列/行里没有现成文字可加。\n'
                '若数据在其他工作表，请在“工作表”下拉里换一张；或确认该列已填内容。'
                % (sheet or '自动'))
            return
        self.set_status('已给 %d 个单元格加字（工作表：%s）' % (n, sheet or '自动'))
        if messagebox.askyesno(
                '完成', '已给工作表“%s”的 %d 个单元格加字：\n%s\n\n是否立即打开查看？'
                % (sheet or '自动', n, tool)):
            self._open(tool)

    def _open(self, path):
        try:
            os.startfile(path)
        except Exception as e:
            messagebox.showwarning('打开失败', '请手动打开：%s\n%s' % (path, e))

    def open_tool(self):
        """打开⑤当前加工/扫描的表格（即加字写入的那个文件）。"""
        tool, _ = self._tool_target()
        if tool and os.path.exists(tool):
            self._open(tool)
        else:
            messagebox.showwarning('提示', '还没有可打开的表格，请先选择。')

    def do_split(self):
        tool, sheet = self._tool_target()
        if not tool or not os.path.exists(tool):
            messagebox.showwarning('提示', '请先选择要加工的表格。')
            return
        try:
            per = int(self.var_per.get().strip())
        except ValueError:
            messagebox.showwarning('提示', '每份行数请填写数字。')
            return
        prefix = os.path.join(self.out_dir, os.path.splitext(os.path.basename(tool))[0])
        try:
            os.makedirs(self.out_dir, exist_ok=True)
        except OSError:
            pass
        try:
            files = xlsx_fill.split_rows(tool, prefix, per, sheet=sheet)
        except (PermissionError, OSError):
            messagebox.showerror('失败', '目录无写权限或表格被占用，请先关闭 WPS/Excel。')
            return
        except Exception as e:
            messagebox.showerror('失败', str(e))
            return
        self.set_status('已切割为 %d 份，全部保存在 %s' % (len(files), self.out_dir))
        listing = '\n'.join(os.path.basename(f) for f in files)
        if len(files) > 8:
            listing = '\n'.join(os.path.basename(f) for f in files[:8]) + \
                '\n…共 %d 份' % len(files)
        if messagebox.askyesno(
                '完成', '已切割为 %d 份，每份都带表头、保留原格式，全部保存在：\n%s\n\n%s'
                '\n\n（切割是多份独立文件，不是一份里只有前 N 行）\n\n是否打开该文件夹查看全部分片？'
                % (len(files), self.out_dir, listing)):
            self._open(self.out_dir)

    def do_extract(self):
        """把含标识（关键字）的行整行抄到新表，格式不变。"""
        tool, sheet = self._tool_target()
        keyword = self.var_mark.get().strip()
        if not tool or not os.path.exists(tool):
            messagebox.showwarning('提示', '请先选择要加工的表格。')
            return
        if not keyword:
            messagebox.showwarning('提示', '请先输入标识（关键字）。')
            return
        safe_kw = re.sub(r'[\\/:*?"<>|]', '_', keyword)
        base = os.path.splitext(os.path.basename(tool))[0]
        try:
            os.makedirs(self.out_dir, exist_ok=True)
        except OSError:
            pass
        out = os.path.join(self.out_dir, '%s-含%s.xlsx' % (base, safe_kw))
        try:
            n = xlsx_fill.extract_rows(tool, out, keyword, sheet=sheet)
        except (PermissionError, OSError):
            messagebox.showerror('失败', '目录无写权限或表格被占用，请先关闭 WPS/Excel。')
            return
        except ValueError as e:
            messagebox.showwarning('没找到', str(e))
            return
        except Exception as e:
            messagebox.showerror('失败', str(e))
            return
        self.set_status('已提取 %d 行' % n)
        if messagebox.askyesno(
                '完成', '含“%s”的行共 %d 行，已抄到：\n%s\n\n是否立即打开？'
                % (keyword, n, out)):
            self._open(out)

    def choose_ytool(self):
        p = filedialog.askopenfilename(title='选择要扫描黄色标识的表格',
                                       filetypes=[('Excel 表格', '*.xlsx'), ('所有文件', '*.*')])
        if p:
            self.var_ytool.set(p)

    def choose_ytarget(self):
        p = filedialog.askopenfilename(title='选择目标表格（黄标行填进它里面）',
                                       filetypes=[('Excel 表格', '*.xlsx'), ('所有文件', '*.*')])
        if p:
            self.var_ytarget.set(p)

    def choose_reddoc(self):
        p = filedialog.askopenfilename(title='选择要标红的 Word 或表格',
                                       filetypes=[('Word/表格', '*.docx *.xlsx'),
                                                  ('Word 文档', '*.docx'),
                                                  ('Excel 表格', '*.xlsx'),
                                                  ('所有文件', '*.*')])
        if p:
            self.var_reddoc.set(p)

    def do_scan_yellow(self):
        """扫描源表格里黄底标识行，确认后按列对应填进用户给的目标表格。"""
        src = self.var_ytool.get().strip()
        tgt = self.var_ytarget.get().strip()
        if not src or not os.path.exists(src):
            messagebox.showwarning('提示', '请先选择带黄色标识的源表格。')
            return
        if not tgt or not os.path.exists(tgt):
            messagebox.showwarning('提示', '请先选择目标表格（黄标行要填进它里面）。')
            return
        try:
            sel = xlsx_fill.marked_rows(src)
        except Exception as e:
            messagebox.showerror('读表失败', str(e))
            return
        if not sel:
            messagebox.showwarning('没找到', '源表格里没扫到带黄色底纹标识的行。')
            return
        # 扫描结果预览：每行取前 3 个非空单元格内容
        lines = []
        for rno, row in sel[:5]:
            vals = [v for _c, v in sorted(row.items(), key=lambda kv: colkey(kv[0]))]
            lines.append('第%d行: %s' % (rno, ' | '.join(vals[:3])))
        if len(sel) > 5:
            lines.append('…… 等共 %d 行' % len(sel))
        if not messagebox.askyesno(
                '扫描结果', '共扫到 %d 行黄色标识：\n%s\n\n要填进目标表格吗？' % (
                    len(sel), '\n'.join(lines))):
            return
        try:
            header = xlsx_fill.read_header(tgt)
            h1 = xlsx_fill.read_header(src)      # 只用表头，别把大表全表解析一遍
        except Exception as e:
            messagebox.showerror('读表失败', str(e))
            return
        # 列对应：目标列名(归一化) → 源表列字母（同名 + 别名）
        src_norm = {norm(v): cl2 for cl2, v in h1.items() if norm(v)}
        col_map = {}
        for cl, text in header.items():
            n = norm(text)
            if not n:
                continue
            if n in src_norm:
                col_map[cl] = src_norm[n]
                continue
            for alias in ALIASES.get(n, []):
                if alias in src_norm:
                    col_map[cl] = src_norm[alias]
                    break
        rows = []
        for _rno, row in sel:
            rv = {}
            for cl in header:
                srccol = col_map.get(cl)          # 源表列字母
                rv[cl] = row.get(srccol, '') if srccol else ''
            rows.append(rv)
        center = bool(self.var_center.get())
        fill = self.var_fill.get().strip() or None
        out = tgt
        try:
            start, end = xlsx_fill.append_rows(tgt, tgt, rows, clear=False,
                                               center=center, fill=fill)
            tail = ''
        except (PermissionError, OSError) as e:
            base, ext = os.path.splitext(tgt)
            fallback = base + '-已填写' + ext
            try:
                start, end = xlsx_fill.append_rows(tgt, fallback, rows, clear=False,
                                                   center=center, fill=fill)
            except Exception as e2:
                messagebox.showerror('填写失败',
                                     '写不进目标表格：%s\n另存副本也失败：%s\n'
                                     '请先关闭正在 WPS/Excel 中打开的表格，再重试。' % (e, e2))
                return
            out = fallback
            tail = '\n注意：目标表格正被打开或无权限，已改存到副本。'
        except Exception as e:
            messagebox.showerror('填写失败', str(e))
            return
        self.set_status('黄标行已填入 %d 行' % len(rows))
        if messagebox.askyesno('完成', '黄色标识行共 %d 行，已填到第 %d~%d 行：\n%s%s\n\n现在打开看看吗？'
                               % (len(rows), start, end, out, tail)):
            os.startfile(out)

    def do_mark_red(self):
        """在 Word 或表格里找到关键字标成红色字体，其它格式不变。"""
        path = self.var_reddoc.get().strip()
        if not path or not os.path.exists(path):
            messagebox.showwarning('提示', '请先在“标红指定内容：文档”一栏选择要标红的 Word 或表格。')
            return
        keyword = self.var_red.get().strip()
        if not keyword:
            messagebox.showwarning('提示', '请先输入要标红的关键字。')
            return
        ext = os.path.splitext(path)[1].lower()
        inplace = bool(self.var_inplace.get())
        out = path if inplace else os.path.join(
            self.out_dir,
            os.path.splitext(os.path.basename(path))[0] + '-标红' + ext)
        try:
            if not inplace:
                os.makedirs(self.out_dir, exist_ok=True)
        except OSError:
            pass
        try:
            if ext == '.docx':
                n = core.mark_red_docx(path, out, keyword)
            elif ext == '.xlsx':
                n = xlsx_fill.mark_red(path, out, keyword)
            else:
                messagebox.showwarning('提示', '只支持 .docx 或 .xlsx 文件。')
                return
        except (PermissionError, OSError):
            messagebox.showerror('失败', '文件正被 WPS/Word/Excel 打开或无写权限，请先关闭再试。')
            return
        except ValueError as e:
            messagebox.showwarning('没找到', str(e))
            return
        except Exception as e:
            messagebox.showerror('失败', str(e))
            return
        self.set_status('已标红 %d 处' % n)
        self.last_red = (path, out)  # 紧接着点“提取标红”可直接用这份结果
        if messagebox.askyesno('完成', '“%s”共标红 %d 处：\n%s\n\n现在打开看看吗？\n（直接点“提取标红→Word”即可提取这份标红结果，无需重新选文件）' % (
                keyword, n, out)):
            os.startfile(out)

    def do_extract_red(self):
        """把标红（红色字体）文字所在的整段/整格提取到一个新的空白 Word。"""
        path = self.var_reddoc.get().strip()
        if not path or not os.path.exists(path):
            messagebox.showwarning('提示', '请先在“标红指定内容：文档”一栏选择已标红的 Word 或表格。')
            return
        ext = os.path.splitext(path)[1].lower()
        stem = os.path.splitext(os.path.basename(path))[0]
        last = getattr(self, 'last_red', None)
        auto = False
        if (last and last[0] == path and os.path.exists(last[1])
                and os.path.splitext(last[1])[1].lower() == ext):
            path, auto = last[1], True  # 刚标完红：自动改用标红结果，不用重新选文件
        out = os.path.join(self.out_dir, stem + '-标红提取.docx')
        try:
            os.makedirs(self.out_dir, exist_ok=True)
        except OSError:
            pass
        try:
            if ext == '.docx':
                n = core.extract_red_docx(path, out)
            elif ext == '.xlsx':
                n = xlsx_fill.extract_red_xlsx(path, out)
            else:
                messagebox.showwarning('提示', '只支持 .docx 或 .xlsx 文件。')
                return
        except PermissionError:
            messagebox.showerror('失败', '文件正被 WPS/Word/Excel 打开或无写权限，请先关闭再试。')
            return
        except ValueError as e:
            messagebox.showwarning('没有标红内容', str(e))
            return
        except Exception as e:
            messagebox.showerror('失败', str(e))
            return
        self.set_status('已提取 %d 段标红内容：%s' % (n, out))
        if messagebox.askyesno('完成', '共提取 %d 段标红内容（整段带走，红色保留）%s：\n%s\n\n现在打开看看吗？' % (
                n, '；自动使用了刚才的标红结果' if auto else '', out)):
            os.startfile(out)

    def choose_notes_src(self):
        ps = filedialog.askopenfilenames(
            title='选择要生成日志/笔记的源文档（可多选）',
            initialdir=self._initial_dir(),
            filetypes=[('图片/表格/Word/文本', '*.png *.jpg *.jpeg *.bmp *.xlsx *.docx *.txt *.md *.log'),
                       ('图片', '*.png *.jpg *.jpeg *.bmp'),
                       ('表格', '*.xlsx'),
                       ('Word 文档', '*.docx'),
                       ('所有文件', '*.*')])
        if ps:
            self.var_notes_src.set(' | '.join(ps))

    def do_make_notes(self):
        """把源文档整理成日志/笔记，存 Word 或记事本。"""
        raw = self.var_notes_src.get().strip()
        if not raw:
            messagebox.showwarning('提示', '请先选择源文档（图片/表格/Word/文本，可多选）。')
            return
        files = [p.strip() for p in raw.split('|') if p.strip()]
        bad = [p for p in files if not os.path.exists(p)]
        if bad:
            messagebox.showwarning('提示', '文件不存在：%s' % bad[0])
            return
        kind = self.var_note_kind.get()
        detailed = self.var_note_detail.get() == '具体'
        try:
            text, img_blocks = doc_notes.build_notes(files, kind=kind, detailed=detailed)
        except ValueError as e:
            messagebox.showwarning('不支持', str(e))
            return
        except Exception as e:
            messagebox.showerror('生成失败', str(e))
            return
        ext = '.docx' if self.var_note_fmt.get() == 'Word' else '.txt'
        out = os.path.join(self.out_dir,
                           '%s_%s%s' % (kind, time.strftime('%Y%m%d-%H%M%S'), ext))
        try:
            os.makedirs(self.out_dir, exist_ok=True)
            if ext == '.docx':
                doc_notes.save_docx(text, out, img_blocks)
            else:
                doc_notes.save_txt(text, out, img_blocks)
        except (PermissionError, OSError) as e:
            messagebox.showerror('保存失败', str(e))
            return
        n_img = sum(len(v) for v in img_blocks.values())
        in_lib = ''
        if self.var_note_tosave.get():
            try:
                shared_notes.upload([out], who=self.var_note_who.get())
                self.refresh_share_lib()
                in_lib = '，已存入共享库（大家可见）'
            except Exception as e:
                in_lib = '，但存入共享库失败：%s' % e
        self.set_status('%s已生成：%s（附 %d 图%s）' % (kind, out, n_img, in_lib))
        if messagebox.askyesno('完成', '%s已生成（共 %d 字%s%s）：\n%s\n\n现在打开看看吗？'
                               % (kind, len(text),
                                  '，随日志携带 %d 张图片' % n_img if n_img else '',
                                  in_lib, out)):
            os.startfile(out)

    def do_save_text_note(self):
        """把文字笔记框里的一段话存成 .txt 放进共享库，大家打开同一共享文件夹就能看到。"""
        text = self.txt_note.get('1.0', 'end').strip()
        if not text:
            messagebox.showwarning('提示', '请先在“文字笔记”框里写点内容。')
            return
        try:
            name = shared_notes.add_text_note(text, who=self.var_note_who.get())
        except Exception as e:
            messagebox.showerror('存入失败', str(e))
            return
        self.txt_note.delete('1.0', 'end')
        self.refresh_share_lib()
        self.set_status('文字笔记已存入共享库：%s' % name)
        messagebox.showinfo('已存入', '文字笔记已存入共享库，大家打开同一共享文件夹就能看到。\n\n文件名：%s' % name)

    # ---------- ⑦ 共享文档库 ----------
    def refresh_share_lib(self):
        """重读共享库目录，列表实时反映别人上传的内容。"""
        if not hasattr(self, 'tree_share'):
            return
        try:
            self.share_items = shared_notes.list_items()
        except Exception as e:
            messagebox.showerror('共享库打不开', str(e))
            return
        self.tree_share.delete(*self.tree_share.get_children())
        for i, (name, who, t, size, _fp) in enumerate(self.share_items):
            self.tree_share.insert('', 'end', iid=str(i), values=(
                name, who, t or '—', '%.0f KB' % (size / 1024.0)))
        self.var_share_dir.set(shared_notes.lib_dir())

    def show_share_usage(self):
        """统计共享库现在占多少硬盘、各人传了多少。"""
        try:
            total, count, rank = shared_notes.usage()
        except Exception as e:
            messagebox.showerror('统计失败', str(e))
            return
        if count == 0:
            messagebox.showinfo('共享库占用',
                                '共享库现在是空的。\n\n位置：%s' % shared_notes.lib_dir())
            return
        lines = ['共 %d 个文件，占用 %s' % (count, shared_notes.fmt_size(total)), '',
                 '按上传人（占用从多到少）：']
        for who, c, s in rank:
            lines.append('  %s：%d 个 · %s' % (who, c, shared_notes.fmt_size(s)))
        lines += ['', '位置：%s' % shared_notes.lib_dir(),
                  '（占的是这台“仓库主机”的硬盘）']
        messagebox.showinfo('共享库占用', '\n'.join(lines))

    def _share_selected(self):
        sel = self.tree_share.selection()
        if not sel:
            messagebox.showwarning('提示', '先在共享库列表里选中一个文件。')
            return []
        return [self.share_items[int(i)] for i in sel]

    def do_upload_share(self):
        """把自己的文档存进共享库，保存后其他人打开工具就能看到。"""
        ps = filedialog.askopenfilenames(
            title='选择要上传到共享库的文档（可多选，任意类型）',
            initialdir=self._initial_dir())
        if not ps:
            return
        try:
            saved = shared_notes.upload(ps, self.var_note_who.get())
        except Exception as e:
            messagebox.showerror('上传失败', str(e))
            return
        self.refresh_share_lib()
        self.set_status('已上传 %d 个文档到共享库' % len(saved))
        if messagebox.askyesno('上传完成', '已保存到共享库（%s）：\n%s\n\n要顺便生成一份日志吗？'
                               % (shared_notes.lib_dir(), '\n'.join(saved))):
            self.var_notes_src.set(' | '.join(
                os.path.join(shared_notes.lib_dir(), n) for n in saved))
            self.do_make_notes()

    def choose_share_dir(self):
        """把共享库换到局域网共享盘等公共目录，换后全库跟着新路径。"""
        d = filedialog.askdirectory(title='选择共享库目录（建议局域网共享盘）')
        if not d:
            return
        try:
            shared_notes.set_dir(d)
        except (ValueError, OSError) as e:
            messagebox.showerror('换路径失败', str(e))
            return
        self.refresh_share_lib()
        messagebox.showinfo('已切换', '共享库现在指向：\n%s\n\n把此路径告诉同事，大家用同一个目录即可互见上传。' % d)

    def open_share_item(self):
        for it in self._share_selected():
            try:
                os.startfile(it[4])
            except Exception as e:
                messagebox.showerror('打开失败', '%s\n%s' % (it[0], e))

    def notes_from_share(self):
        """选中共享库文件 → 填进源文档并直接出日志/笔记。"""
        its = self._share_selected()
        if not its:
            return
        self.var_notes_src.set(' | '.join(it[4] for it in its))
        self.do_make_notes()

    def delete_share_item(self):
        its = self._share_selected()
        if not its:
            return
        if not messagebox.askyesno('确认删除', '从共享库删除选中的 %d 个文件？\n（只删库里的副本，不动上传者本机原文件）' % len(its)):
            return
        for it in its:
            shared_notes.delete(it[0])
        self.refresh_share_lib()

    def choose_demo_file(self):
        p = filedialog.askopenfilename(
            title='选择参数文件（Word/表格/文本）',
            filetypes=[('Word/表格/文本', '*.docx *.doc *.wps *.xlsx *.txt *.md'),
                       ('Word 文档', '*.docx *.doc *.wps'), ('Excel 表格', '*.xlsx'),
                       ('所有文件', '*.*')])
        if not p:
            return
        self.var_demo_file.set(p)
        try:
            ext = os.path.splitext(p)[1].lower()
            if ext in ('.docx', '.doc', '.wps'):
                from docx import Document
                doc = Document(core.to_docx(p))
                lines = [t.text.strip() for t in doc.paragraphs if t.text.strip()]
                for tb in doc.tables:
                    for row in tb.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            lines.append('、'.join(cells))
            elif ext == '.xlsx':
                _h, rows = xlsx_fill.read_rows(p)
                lines = []
                for row in rows:
                    vals = [row[c] for c in sorted(row, key=colkey) if row.get(c)]
                    if vals:
                        lines.append('、'.join(vals))
            else:
                raw = open(p, 'rb').read()
                for enc in ('utf-8-sig', 'utf-8', 'gbk'):
                    try:
                        lines = raw.decode(enc).splitlines()
                        break
                    except UnicodeDecodeError:
                        continue
                else:
                    lines = raw.decode('gbk', errors='replace').splitlines()
        except Exception as e:
            messagebox.showerror('读取失败', str(e))
            return
        self.txt_demo.delete('1.0', 'end')
        self.txt_demo.insert('1.0', '\n'.join(lines))
        self.set_status('已载入参数文件，共 %d 行，可编辑后点生成' % len(lines))

    def do_make_demo(self):
        """参数内容 → 单文件可交互投标演示页面。"""
        text = self.txt_demo.get('1.0', 'end').strip()
        if not text:
            messagebox.showwarning('提示', '请先粘贴参数内容，或浏览载入参数文件。')
            return
        try:
            result = demo_page.generate_demo(text, self.out_dir, self.demo_history_path)
        except Exception as e:
            messagebox.showerror('生成失败', '本次未计入生成次数。\n%s' % e)
            return
        out = result['path']
        self.set_status('方案 %d/3 · %s：%s' % (result['variant'], result['scheme'], out))
        if messagebox.askyesno('完成',
                '已生成方案 %d/3（同一内容第 %d 次）：%s\n%d 个模块、%d 项功能、%d 个页面、%d 个导航入口\n'
                '前三次依次给蓝/紫/绿三套，之后循环复用三套风格继续生成；文件不会互相覆盖。\n%s\n\n现在打开看看吗？' %
                (result['variant'], result['attempt'], result['scheme'], result['modules'], result['features'],
                 result['pages'], result['navigation'], out)):
            os.startfile(out)

    def set_status(self, text, warn=False):
        self.lbl_status.config(text=text, foreground=('#c60' if warn else '#0a7'))

    # ---------------- ⑩ 生成PPT ----------------
    def choose_ppt_file(self):
        p = filedialog.askopenfilename(
            title='选择生成PPT的来源文件（Word/表格/文本）',
            filetypes=[('Word/表格/文本', '*.docx *.doc *.wps *.xlsx *.txt *.md'),
                       ('Word 文档', '*.docx *.doc *.wps'), ('所有文件', '*.*')])
        if not p:
            return
        self.var_ppt_file.set(p)
        try:
            text = ppt_gen.extract_source(p)
        except Exception as e:
            messagebox.showerror('载入失败', '%s\n%s' % (p, e))
            return
        self.txt_ppt.delete('1.0', 'end')
        self.txt_ppt.insert('1.0', text)
        if not self.var_ppt_title.get().strip():
            self.var_ppt_title.set(os.path.splitext(os.path.basename(p))[0])
        self.set_status('已载入 %d 字到内容框，可直接编辑后点“生成PPT”' % len(text))

    def do_make_ppt(self):
        """内容框文本（或已载入的 Word/描述）→ 可放映的 .pptx。"""
        text = self.txt_ppt.get('1.0', 'end').strip()
        if not text:
            messagebox.showwarning('提示', '请先浏览 Word 文档，或在内容框里填写描述。')
            return
        name = self.var_ppt_scheme.get()
        idx = None if name == '自动轮流' else \
            [s['name'] for s in ppt_gen.SCHEMES].index(name)
        try:
            pages = int(self.var_ppt_pages.get().strip() or 0)
        except ValueError:
            pages = 0
        if pages < 1 or pages > 30:
            messagebox.showwarning('提示', '页数请填 1~30 之间的整数。')
            return
        try:
            result = ppt_gen.generate_ppt(text, self.out_dir,
                                          self.var_ppt_title.get(), idx, pages)
        except Exception as e:
            messagebox.showerror('生成失败', str(e))
            return
        out = result['path']
        self.set_status('已生成“%s”· %s · 内容%d页：%s' % (
            result['deck'], result['scheme'], result['slides'], out))
        if messagebox.askyesno('完成',
                'PPT 已生成：%s\n配色：%s · 内容 %d 页 / %d 条要点（另含封面、目录、结束页）\n%s\n\n现在打开看看吗？' % (
                result['deck'], result['scheme'], result['slides'], result['bullets'], out)):
            os.startfile(out)

    # ---------------- ⑪-⑯ 办公工具箱 ----------------
    def _frow(self, parent, label, var, cmd, hint=''):
        fr = ttk.Frame(parent); fr.pack(fill='x', padx=6, pady=2)
        ttk.Label(fr, text=label, width=14).pack(side='left')
        ttk.Entry(fr, textvariable=var, width=30).pack(side='left', padx=4)
        ttk.Button(fr, text='浏览…', command=cmd).pack(side='left')
        if hint:
            ttk.Label(fr, text=hint).pack(side='left', padx=4)

    def _multi(self, var):
        return [p.strip() for p in var.get().split('|') if p.strip()]

    def _pick_one(self, var, title, types=None):
        p = filedialog.askopenfilename(title=title, initialdir=self._initial_dir(),
                                       filetypes=types or [('所有文件', '*.*')])
        if p:
            var.set(p)

    # ⑪ 合并
    def choose_merge_files(self):
        ps = filedialog.askopenfilenames(title='选择要合并的 Excel（可多选，表头需一致）',
                                         initialdir=self._initial_dir(),
                                         filetypes=[('Excel 表格', '*.xlsx')])
        if ps:
            self.var_merge_files.set(' | '.join(ps))

    def do_merge(self):
        files = self._multi(self.var_merge_files)
        if len(files) < 2:
            messagebox.showwarning('提示', '请至少选择 2 份表格。')
            return
        out = os.path.join(self.out_dir, '合并汇总_%s.xlsx' % time.strftime('%Y%m%d-%H%M%S'))
        try:
            total, per, dup = office_tools.merge_xlsx(
                files, out, dedup=self.var_merge_dedup.get())
        except Exception as e:
            messagebox.showerror('合并失败', str(e))
            return
        detail = '\n'.join('  %s：%d 行' % (n, c) for n, c in per)
        self.set_status('已合并 %d 份 → %d 行（去重%d）：%s' % (len(files), total, dup, out))
        if messagebox.askyesno('完成', '共合并 %d 行（来自 %d 份，去重删除 %d 行）：\n%s\n\n保存到：\n%s\n\n现在打开看看吗？'
                               % (total, len(files), dup, detail, out)):
            self._open(out)

    # ⑫ 套打
    def choose_mm_tpl(self):
        self._pick_one(self.var_mm_tpl, '选择模板文档',
                       [('Word 文档', '*.docx *.doc *.wps')])

    def choose_mm_data(self):
        self._pick_one(self.var_mm_data, '选择名单表格', [('Excel 表格', '*.xlsx')])

    def do_mailmerge(self):
        tpl = self.var_mm_tpl.get().strip(); data = self.var_mm_data.get().strip()
        if not (tpl and data):
            messagebox.showwarning('提示', '请选择模板文档和名单表格。')
            return
        outdir = os.path.join(self.out_dir, '套打_%s' % time.strftime('%Y%m%d-%H%M%S'))
        try:
            outs = office_tools.mail_merge(tpl, data, outdir,
                                           name_field=self.var_mm_name.get().strip() or None)
        except Exception as e:
            messagebox.showerror('生成失败', str(e))
            return
        self.set_status('已批量套打 %d 份 → %s' % (len(outs), outdir))
        if messagebox.askyesno('完成', '已生成 %d 份文档，保存在：\n%s\n\n打开该文件夹看看吗？'
                               % (len(outs), outdir)):
            self._open(outdir)

    # ⑬ 对账
    def choose_rec_a(self):
        self._pick_one(self.var_rec_a, '选择表A', [('Excel 表格', '*.xlsx')])

    def choose_rec_b(self):
        self._pick_one(self.var_rec_b, '选择表B', [('Excel 表格', '*.xlsx')])

    def do_reconcile(self):
        a = self.var_rec_a.get().strip(); b = self.var_rec_b.get().strip()
        key = self.var_rec_key.get().strip()
        if not (a and b and key):
            messagebox.showwarning('提示', '请选择两张表并填写关键列名。')
            return
        rep = os.path.join(self.out_dir, '对账报告_%s.txt' % time.strftime('%Y%m%d-%H%M%S'))
        try:
            r = office_tools.reconcile(a, b, key, report_out=rep)
        except Exception as e:
            messagebox.showerror('对账失败', str(e))
            return
        msg = ('表A %d 条 / 表B %d 条\n两表都有：%d\n只在表A：%d\n只在表B：%d\n同键但值不同：%d\n\n报告：%s'
               % (r['total_a'], r['total_b'], r['both'], len(r['only_a']),
                  len(r['only_b']), len(r['diffs']), rep))
        self.set_status('对账完成：差异 %d 项'
                        % (len(r['only_a']) + len(r['only_b']) + len(r['diffs'])))
        if messagebox.askyesno('对账结果', msg + '\n\n打开报告看看吗？'):
            self._open(rep)

    # ⑭ 清洗
    def choose_clean_file(self):
        self._pick_one(self.var_clean_file, '选择要清洗的表格', [('Excel 表格', '*.xlsx')])

    def do_clean(self):
        f = self.var_clean_file.get().strip()
        if not f:
            messagebox.showwarning('提示', '请选择要清洗的表格。')
            return
        out = os.path.join(self.out_dir,
                           os.path.splitext(os.path.basename(f))[0] + '-已清洗.xlsx')
        try:
            n = office_tools.clean_xlsx(f, out)
        except Exception as e:
            messagebox.showerror('清洗失败', str(e))
            return
        self.set_status('已清洗 %d 个单元格 → %s' % (n, out))
        if messagebox.askyesno('完成', '已清洗 %d 处（去空格/全角转半角），另存为：\n%s\n\n打开看看吗？'
                               % (n, out)):
            self._open(out)

    # ⑮ 转换 + 重命名
    def choose_conv_files(self):
        ps = filedialog.askopenfilenames(title='选择要转换的文件', initialdir=self._initial_dir(),
                                         filetypes=[('Word/Excel', '*.doc *.wps *.xlsx'),
                                                    ('所有文件', '*.*')])
        if ps:
            self.var_conv_files.set(' | '.join(ps))

    def do_convert(self):
        files = self._multi(self.var_conv_files)
        if not files:
            messagebox.showwarning('提示', '请选择要转换的文件。')
            return
        done, errs = [], []
        for f in files:
            ext = os.path.splitext(f)[1].lower()
            try:
                if ext in ('.doc', '.wps'):
                    done.append(office_tools.convert_to_docx(f, self.out_dir))
                elif ext == '.xlsx':
                    p, _ = office_tools.xlsx_to_csv(
                        f, os.path.join(self.out_dir,
                                        os.path.splitext(os.path.basename(f))[0] + '.csv'))
                    done.append(p)
                else:
                    errs.append('%s（不支持）' % os.path.basename(f))
            except Exception as e:
                errs.append('%s：%s' % (os.path.basename(f), e))
        self.set_status('转换完成 %d 个' % len(done))
        messagebox.showinfo('转换完成',
                            '成功 %d 个，存到：%s\n\n%s%s' % (
                                len(done), self.out_dir,
                                '\n'.join(os.path.basename(d) for d in done),
                                ('\n\n失败：\n' + '\n'.join(errs)) if errs else ''))
        if done:
            self._open(self.out_dir)

    def choose_ren_dir(self):
        d = filedialog.askdirectory(title='选择要重命名的文件夹', initialdir=self._initial_dir())
        if d:
            self.var_ren_dir.set(d)

    def do_rename(self):
        folder = self.var_ren_dir.get().strip()
        if not folder:
            messagebox.showwarning('提示', '请选择文件夹。')
            return
        try:
            plan = office_tools.rename_plan(
                folder, self.var_ren_mode.get(), find=self.var_ren_find.get(),
                repl=self.var_ren_repl.get(), prefix=self.var_ren_prefix.get(),
                ext_filter=(self.var_ren_ext.get().strip() or None))
        except Exception as e:
            messagebox.showerror('无法重命名', str(e))
            return
        if not plan:
            messagebox.showinfo('没有变化', '没有需要重命名的文件。')
            return
        preview = '\n'.join('%s  →  %s' % (o, n) for _op, _np, o, n in plan[:15])
        more = '' if len(plan) <= 15 else '\n…共 %d 个' % len(plan)
        if not messagebox.askyesno('确认重命名', '将对 %d 个文件重命名：\n%s%s\n\n确定执行吗？'
                                    % (len(plan), preview, more)):
            return
        n = office_tools.apply_rename(plan)
        self.set_status('已重命名 %d 个文件' % n)
        messagebox.showinfo('完成', '已重命名 %d 个文件。' % n)

    # ⑯ 周报/纪要
    def do_weekly(self):
        text = self.txt_wk.get('1.0', 'end').strip()
        if not text:
            messagebox.showwarning('提示', '请先贴入要点（每行一条）。')
            return
        kind = self.var_wk_kind.get()
        who = self.var_wk_who.get().strip() or shared_notes.user()
        lines = [l.strip() for l in text.splitlines() if l.strip()]
        now = time.strftime('%Y-%m-%d %H:%M')
        if kind == '周报':
            body = ['工作周报', '', '记录人：%s    日期：%s' % (who, now), '', '一、本周工作']
            body += ['%d. %s' % (i, t) for i, t in enumerate(lines, 1)]
            body += ['', '二、下周计划', '（待补充）']
        else:
            body = ['会议纪要', '', '时间：%s    记录人：%s' % (now, who), '', '一、议题与要点']
            body += ['· %s' % t for t in lines]
            body += ['', '二、结论与待办', '（待补充）']
        out = os.path.join(self.out_dir, '%s_%s.docx' % (kind, time.strftime('%Y%m%d-%H%M%S')))
        try:
            doc_notes.save_docx('\n'.join(body), out)
            shared_notes.upload([out], who=who)
        except Exception as e:
            messagebox.showerror('生成失败', str(e))
            return
        self.refresh_share_lib()
        self.set_status('%s已生成并存入共享库：%s' % (kind, out))
        if messagebox.askyesno('完成', '%s已生成并存入共享库（大家可见）：\n%s\n\n现在打开看看吗？'
                               % (kind, out)):
            self._open(out)

    # ---------------- ⬓ 视频合成 ----------------
    def _refresh_video_list(self):
        self.lst_video.delete(0, 'end')
        for p in self.video_files:
            self.lst_video.insert('end', os.path.basename(p))

    def add_videos(self):
        if len(self.video_files) >= video_tools.MAX_VIDEOS:
            messagebox.showwarning('提示', '最多只能添加 %d 个视频。' % video_tools.MAX_VIDEOS)
            return
        ps = filedialog.askopenfilenames(
            title='选择视频（可多选，最多 %d 个）' % video_tools.MAX_VIDEOS,
            initialdir=self._initial_dir(),
            filetypes=[('视频文件', '*.mp4 *.mov *.mkv *.avi *.flv *.wmv *.m4v *.webm *.ts *.mpg *.mpeg *.3gp'),
                       ('所有文件', '*.*')])
        for p in ps:
            if p in self.video_files:
                continue
            if len(self.video_files) >= video_tools.MAX_VIDEOS:
                break
            self.video_files.append(p)
        self._refresh_video_list()

    def remove_video(self):
        for i in reversed(list(self.lst_video.curselection())):
            self.video_files.pop(i)
        self._refresh_video_list()

    def move_video(self, step):
        sel = self.lst_video.curselection()
        if not sel:
            return
        i = sel[0]; j = i + step
        if 0 <= j < len(self.video_files):
            self.video_files[i], self.video_files[j] = self.video_files[j], self.video_files[i]
            self._refresh_video_list()
            self.lst_video.selection_clear(0, 'end')
            self.lst_video.selection_set(j)

    def clear_videos(self):
        self.video_files = []
        self._refresh_video_list()

    def do_video_merge(self, mode='original'):
        if len(self.video_files) < 2:
            messagebox.showwarning('提示', '请至少添加 2 个视频。')
            return
        if video_tools.find_ffmpeg() is None:
            messagebox.showerror(
                '缺视频引擎',
                '未找到 ffmpeg，无法合成视频。\n\n请用内置 ffmpeg 的“填表工具.exe”，\n'
                '或把 ffmpeg.exe 放到本程序目录的 tools 文件夹后重试。')
            return
        name = self.var_video_out.get().strip() or '合成视频.mp4'
        if not name.lower().endswith('.mp4'):
            name += '.mp4'
        out = os.path.join(self.out_dir, name)
        files = list(self.video_files)
        # 画质：原画质=无损优先直拼；高清=强制高画质重编码
        if mode == 'hd':
            crf, preset, allow_lossless = 16, 'medium', False
        else:
            crf, preset, allow_lossless = 18, 'veryfast', True
        target_h = {'放大到1080P': 1080, '放大到2K': 1440,
                    '放大到4K(慢)': 2160}.get(self.var_video_res.get())
        self.btn_video.config(state='disabled')
        self.btn_video_hd.config(state='disabled')
        self.lbl_video.config(text='合成中…')
        self.set_status('视频合成中（%d 段·%s）…'
                        % (len(files), '高清' if mode == 'hd' else '原画质'))

        def work():
            try:
                res = video_tools.concat_videos(
                    files, out, crf=crf, preset=preset,
                    target_h=target_h, allow_lossless=allow_lossless,
                    progress_cb=lambda i, t, m: self.root.after(
                        0, lambda m=m, i=i, t=t: self.lbl_video.config(
                            text='%d/%d %s' % (i, t, m))))
            except Exception as e:
                self.root.after(0, lambda e=str(e): self._video_done(None, e, len(files)))
                return
            self.root.after(0, lambda r=res: self._video_done(r, None, len(files)))

        threading.Thread(target=work, daemon=True).start()

    def _video_done(self, res, err, count):
        self.btn_video.config(state='normal')
        self.btn_video_hd.config(state='normal')
        if err:
            self.lbl_video.config(text='')
            self.set_status('视频合成失败')
            messagebox.showerror('合成失败', err)
            return
        self.lbl_video.config(text='已完成')
        self.set_status('视频已合成：%s' % res)
        if messagebox.askyesno('完成', '已把 %d 个视频合成为一个：\n%s\n\n现在打开看看吗？'
                               % (count, res)):
            self._open(res)

    # ---------------- ⬔ 去水印 ----------------
    def _wm_set_path(self, p):
        self.wm_path = p
        self.wm_box = None
        self.lbl_wm_box.config(text='未框选区域', foreground='#888')
        kind = '照片' if watermark_tools.is_image(p) else '视频'
        self.lbl_wm_src.config(text='已选：%s（%s）' % (os.path.basename(p), kind),
                               foreground='#333')

    def choose_wm_file(self):
        p = filedialog.askopenfilename(
            title='选择要去水印的照片或视频', initialdir=self._initial_dir(),
            filetypes=[('图片/视频',
                        '*.png *.jpg *.jpeg *.bmp *.webp *.tif *.tiff '
                        '*.mp4 *.mov *.mkv *.avi *.flv *.wmv *.m4v *.webm *.ts *.mpg *.mpeg'),
                       ('所有文件', '*.*')])
        if p:
            self.var_wm_src.set(p)
            self._wm_set_path(p)

    def wm_use_url(self):
        s = self.var_wm_src.get().strip()
        if not s:
            messagebox.showwarning('提示', '请先在框里粘贴网址（http/https）。')
            return
        if os.path.exists(s):
            self._wm_set_path(s)
            return
        self.lbl_wm_src.config(text='下载中…', foreground='#08979c')
        self.root.update_idletasks()
        try:
            p = watermark_tools.download(s, self.out_dir)
        except Exception as e:
            self.lbl_wm_src.config(text='下载失败', foreground='#c0392b')
            messagebox.showerror('下载失败', str(e))
            return
        self.var_wm_src.set(p)
        self._wm_set_path(p)

    def wm_pick_box(self):
        if not self.wm_path or not os.path.exists(self.wm_path):
            messagebox.showwarning('提示', '请先选文件或用网址下载。')
            return
        try:
            if watermark_tools.is_image(self.wm_path):
                preview = self.wm_path
                size = watermark_tools.image_size(self.wm_path)
            else:
                preview = os.path.join(tempfile.gettempdir(),
                                       'wm_frame_%d.png' % int(time.time() * 1000))
                watermark_tools.grab_frame(self.wm_path, preview)
                size = watermark_tools.image_size(preview)
        except Exception as e:
            messagebox.showerror('无法打开画面', str(e))
            return
        WatermarkPicker(self.root, preview, size, self._wm_box_chosen)

    def _wm_box_chosen(self, box):
        self.wm_box = box
        self.lbl_wm_box.config(text='已框选：x=%d y=%d 宽%d 高%d' % box, foreground='#0877b6')

    def do_watermark(self):
        if not self.wm_path or not os.path.exists(self.wm_path):
            messagebox.showwarning('提示', '请先选文件或下载网址。')
            return
        if not self.wm_box:
            messagebox.showwarning('提示', '请先点“① 框选水印位置”圈出水印。')
            return
        is_img = watermark_tools.is_image(self.wm_path)
        ext = os.path.splitext(self.wm_path)[1] or ('.png' if is_img else '.mp4')
        name = self.var_wm_out.get().strip() or '去水印结果'
        out = os.path.join(self.out_dir, name + ext)
        path = self.wm_path
        box = self.wm_box
        method = 'auto' if self.var_wm_method.get() == '智能修补' else 'delogo'
        self.btn_wm.config(state='disabled')
        self.lbl_wm.config(text='处理中…')
        self.set_status('去水印处理中…')

        def work():
            try:
                if is_img:
                    res = watermark_tools.remove_image(path, out, box, method=method)
                else:
                    res = watermark_tools.remove_video(
                        path, out, box,
                        progress_cb=lambda m: self.root.after(
                            0, lambda m=m: self.lbl_wm.config(text=m)))
            except Exception as e:
                self.root.after(0, lambda e=str(e): self._wm_done(None, e))
                return
            self.root.after(0, lambda r=res: self._wm_done(r, None))

        threading.Thread(target=work, daemon=True).start()

    def _wm_done(self, res, err):
        self.btn_wm.config(state='normal')
        if err:
            self.lbl_wm.config(text='')
            self.set_status('去水印失败')
            messagebox.showerror('去除失败', err)
            return
        self.lbl_wm.config(text='已完成')
        self.set_status('已去除水印：%s' % res)
        if messagebox.askyesno('完成', '已去除水印并保存到：\n%s\n\n现在打开看看吗？' % res):
            self._open(res)


class WatermarkPicker(tk.Toplevel):
    """弹窗显示图片/视频帧，鼠标拖框选水印区域，确认回调原图像素坐标 (x,y,w,h)。"""
    MAXW, MAXH = 860, 520

    def __init__(self, parent, image_path, orig_size, on_pick):
        super().__init__(parent)
        self.title('拖拽方框圈出水印位置，然后点“确认此区域”')
        self.orig_w, self.orig_h = orig_size
        self.on_pick = on_pick
        self._box = None
        self._start = None
        self._rect = None
        from PIL import Image, ImageTk
        self.scale = min(self.MAXW / self.orig_w, self.MAXH / self.orig_h)
        disp_w = max(1, int(self.orig_w * self.scale))
        disp_h = max(1, int(self.orig_h * self.scale))
        im = Image.open(image_path).convert('RGB').resize((disp_w, disp_h), Image.LANCZOS)
        self._photo = ImageTk.PhotoImage(im)
        self.canvas = tk.Canvas(self, width=disp_w, height=disp_h, cursor='cross',
                                highlightthickness=0)
        self.canvas.pack(padx=8, pady=8)
        self.canvas.create_image(0, 0, anchor='nw', image=self._photo)
        self.canvas.bind('<ButtonPress-1>', self._down)
        self.canvas.bind('<B1-Motion>', self._drag)
        self.canvas.bind('<ButtonRelease-1>', self._up)
        bar = ttk.Frame(self); bar.pack(fill='x', padx=8, pady=(0, 8))
        self.lbl = ttk.Label(bar, text='按住鼠标左键拖出一个方框')
        self.lbl.pack(side='left')
        ttk.Button(bar, text='确认此区域', command=self._confirm).pack(side='right', padx=4)
        ttk.Button(bar, text='取消', command=self.destroy).pack(side='right')
        self.transient(parent)
        self.grab_set()

    def _to_orig(self, cx, cy):
        x = max(0, min(int(cx / self.scale), self.orig_w))
        y = max(0, min(int(cy / self.scale), self.orig_h))
        return x, y

    def _down(self, e):
        self._start = (e.x, e.y)
        self._rect = self.canvas.create_rectangle(e.x, e.y, e.x, e.y, outline='red', width=2)

    def _drag(self, e):
        if self._start and self._rect:
            self.canvas.coords(self._rect, self._start[0], self._start[1], e.x, e.y)

    def _up(self, e):
        if not self._start:
            return
        x0, y0 = self._to_orig(*self._start)
        x1, y1 = self._to_orig(e.x, e.y)
        bx, by = min(x0, x1), min(y0, y1)
        bw, bh = abs(x1 - x0), abs(y1 - y0)
        if bw >= 2 and bh >= 2:
            self._box = (bx, by, bw, bh)
            self.lbl.config(text='已选区域：x=%d y=%d 宽%d 高%d（可重拖调整）' % self._box)
        self._start = None

    def _confirm(self):
        if not self._box:
            self.lbl.config(text='还没有框，请先拖出一个方框')
            return
        try:
            self.on_pick(self._box)
        finally:
            self.destroy()


def main():
    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == '__main__':
    main()
