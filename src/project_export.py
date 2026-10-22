# -*- coding: utf-8 -*-
"""新文书安全导出：仅复制样式特征，正文/金额来自已校验共享模型。"""
from copy import copy
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import os
import re
import tempfile
import uuid

from project_docs import KINDS, shared_tables, validate, number
from project_sources import word_style, excel_style


def _safe_name(value):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(value)).strip(' .')[:70]
    if not name or re.fullmatch(r'(?i)(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])', name):
        name = '项目文书'
    return name


def _clean(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))


def _chapter_name(index, title):
    if re.match(r'^(?:[一二三四五六七八九十]+[、．.]|\d+[、．.])', title):
        return title
    digits = '零一二三四五六七八九'
    if index < 10:
        prefix = digits[index]
    else:
        tens, ones = divmod(index, 10)
        prefix = (digits[tens] if tens > 1 else '') + '十' + (digits[ones] if ones else '')
    return prefix + '、' + title


def _font(run, family, size, bold=False, color=None):
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    run.font.name = family
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), family)
    run.font.size = Pt(size)
    run.bold = bold
    if color and re.fullmatch('[0-9a-fA-F]{6}', color):
        run.font.color.rgb = RGBColor.from_string(color)


def _paragraph(doc, text, profile, heading=None):
    from docx.shared import Pt
    p = doc.add_paragraph(style='Heading %d' % heading if heading else None)
    body = profile.get('heading%d' % heading, {}) if heading else profile.get('body', {})
    r = p.add_run(_clean(text))
    _font(r, body.get('font') or ('黑体' if heading else '宋体'),
          body.get('size') or ({1: 17, 2: 14, 3: 12}.get(heading, 12)),
          body.get('bold', bool(heading)), body.get('color') or '000000')
    fmt = p.paragraph_format
    fmt.line_spacing = Pt(body['line_spacing_points']) if 'line_spacing_points' in body else body.get('line_spacing', 1.5)
    fmt.space_before = Pt(body.get('space_before', 10 if heading else 0))
    fmt.space_after = Pt(body.get('space_after', 6))
    fmt.first_line_indent = Pt(body.get('first_line_indent', 0 if heading else 24))
    fmt.widow_control = True
    if heading:
        fmt.keep_with_next = True
    if body.get('alignment') is not None:
        p.alignment = body['alignment']
    return p


def _word_table(doc, table, profile):
    from docx.enum.section import WD_SECTION_START, WD_ORIENT
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    wide = len(table['headers']) >= 7
    before = doc.sections[-1]
    portrait = (before.page_width, before.page_height)
    if wide and before.page_width < before.page_height:
        section = doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = portrait[1], portrait[0]
    _paragraph(doc, table['title'], profile, 3)
    result = doc.add_table(rows=1, cols=len(table['headers']))
    result.style = 'Table Grid'
    result.autofit = False
    usable = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
    weights = [3 if any(w in h for w in ('规格', '口径', '原文', '链接')) else
               2 if any(w in h for w in ('内容', '名称', '成果')) else 1 for h in table['headers']]
    for col, weight in zip(result.columns, weights):
        col.width = int(usable * weight / sum(weights))
    repeat = OxmlElement('w:tblHeader')
    result.rows[0]._tr.get_or_add_trPr().append(repeat)
    for rno, values in enumerate([table['headers']] + table['rows']):
        cells = result.rows[0].cells if rno == 0 else result.add_row().cells
        for column, (cell, value) in enumerate(zip(cells, values)):
            cell.width = int(usable * weights[column] / sum(weights))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(_clean(value))
            style = profile.get('table_header', {}) if rno == 0 else profile.get('body', {})
            _font(run, style.get('font', '宋体'), min(style.get('size', 10.5), 12), rno == 0, style.get('color') or '000000')
            if rno == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), profile.get('header_fill', 'DFEAF4'))
                cell._tc.get_or_add_tcPr().append(shade)
            if profile.get('border'):
                borders = OxmlElement('w:tcBorders')
                for side in ('top', 'left', 'bottom', 'right'):
                    element = OxmlElement('w:' + side)
                    for key, val in profile['border'].items():
                        if val:
                            element.set(qn('w:' + key), val)
                    borders.append(element)
                cell._tc.get_or_add_tcPr().append(borders)
    if wide and portrait[0] < portrait[1]:
        section = doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = portrait


def _citations_table(project):
    rows = []
    seen = set()
    for fact in project['facts']:
        for cite in fact['citations']:
            if (cite['source_id'], cite['quote']) in seen:
                continue
            seen.add((cite['source_id'], cite['quote']))
            block = next((b for b in project['sources']['blocks'] if b['id'] == cite['source_id']), {})
            rows.append([cite['source_id'], block.get('file', ''), block.get('location', ''),
                         cite['quote'], block.get('published', ''), block.get('retrieved', '')])
    return {'title': '编制依据核对表（来源记载不等于独立证实）',
            'headers': ['来源ID', '资料名称', '原文位置/链接', '原文片段', '发布日期', '获取时间'], 'rows': rows}


def write_word(project, path, style=None, include_notes=False):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    profile, _ = word_style(style)
    normal = doc.styles['Normal']
    body = profile.get('body', {})
    normal.font.name = body.get('font') or '宋体'
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), normal.font.name)
    normal.font.size = Pt(body.get('size') or 12)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.3)
    section.left_margin = section.right_margin = Cm(2.5)
    for key, value in profile.get('page', {}).items():
        if value > 0:
            setattr(section, key, Pt(value))
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(48)
    _font(title.add_run(project['title']), '黑体', 23, True)
    p = doc.add_paragraph(project['kind'])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('申报单位：' + str(project['model'].get('school') or '【待补充：申报单位】'))
    doc.add_paragraph('编制日期：' + datetime.now().strftime('%Y年%m月%d日'))
    _paragraph(doc, '实施时间口径：' + str(project['model'].get('time_basis') or '【待补充：起止点及采购是否计入】'), profile)
    if project['issues']:
        doc.add_paragraph('编制草稿 · 含待确认事项；请复核来源、数据、目标与政策适用性后使用。')
    doc.add_page_break()
    _paragraph(doc, '目录', profile, 1)
    for n, section_data in enumerate(sorted(project['sections'], key=lambda s: s['index']), 1):
        p = doc.add_paragraph()
        link = OxmlElement('w:hyperlink')
        link.set(qn('w:anchor'), 'chapter_%d' % n)
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = _chapter_name(n, section_data['title'])
        run.append(text)
        link.append(run)
        p._p.append(link)
    doc.add_paragraph('点击章节名称跳转。目录不预填未经排版计算的页码。')
    doc.add_page_break()
    tables = shared_tables(project)
    used = set()
    for n, section_data in enumerate(sorted(project['sections'], key=lambda s: s['index']), 1):
        heading = _paragraph(doc, _chapter_name(n, section_data['title']), profile, 1)
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(n))
        start.set(qn('w:name'), 'chapter_%d' % n)
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), str(n))
        heading._p.insert(0, start)
        heading._p.append(end)
        for block in section_data['blocks']:
            if block.get('heading'):
                _paragraph(doc, block['heading'], profile, 2)
            for paragraph in block['paragraphs']:
                _paragraph(doc, paragraph, profile)
            for ref in block.get('table_refs', []):
                if ref in tables:
                    _word_table(doc, tables[ref], profile)
                    used.add(ref)
    required = {'performance', 'investment', 'funding'} if project['kind'] == KINDS[3] else set(tables)
    missing = required - used
    if missing:
        _paragraph(doc, '配套表格', profile, 1)
        for ref, table in tables.items():
            if ref in missing:
                _word_table(doc, table, profile)
    if include_notes:
        _paragraph(doc, '编制核对事项', profile, 1)
        for issue in project['issues']:
            _paragraph(doc, '[%s] %s' % (issue['level'], issue['message']), profile)
        _word_table(doc, _citations_table(project), profile)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
    footer.add_run(' 页')
    doc.core_properties.author = ''
    doc.core_properties.title = project['title']
    doc.save(path)


def _cell_value(cell, value, numeric=False):
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    value = ILLEGAL_CHARACTERS_RE.sub('', _clean(value))
    if len(value) > 32767:
        raise ValueError('单个Excel单元格超过32767字符，请拆分该段正文。')
    parsed = number(value) if numeric else None
    if parsed is not None:
        cell.value = float(parsed)
        cell.number_format = '0.00' if parsed != parsed.to_integral_value() else '0'
    else:
        cell.value = value
        cell.data_type = 's'


def write_excel(project, path, style=None, include_notes=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    profile, _ = excel_style(style)
    thin = Side(style='thin', color='C7D2DE')
    def sheet(title, headers, rows):
        name = re.sub(r'[\\/*?:\[\]]', '_', title)[:31]
        ws = wb.create_sheet(name)
        for rno, row in enumerate([headers] + rows, 1):
            for cno, value in enumerate(row, 1):
                cell = ws.cell(rno, cno)
                header = headers[cno - 1] if cno <= len(headers) else ''
                numeric = rno > 1 and any(word in header for word in ('金额', '单价', '数量', '成本'))
                _cell_value(cell, value, numeric)
                cell.font = Font(name='宋体', size=11, bold=rno == 1, color='17324D')
                cell.fill = PatternFill('solid', fgColor='DFEAF4' if rno == 1 else 'FFFFFF')
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                for key, item in profile.get('header' if rno == 1 else 'body', {}).items():
                    setattr(cell, key, copy(item))
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                cell.alignment = alignment
        for cno, header in enumerate(headers, 1):
            letter = get_column_letter(cno)
            default = 48 if any(w in header for w in ('正文', '片段', '范围', '口径', '说明', '工作内容', '链接')) else 20
            if len(headers) >= 7:
                default = 30 if default == 48 else 22 if any(w in header for w in ('内容', '名称')) else 14
            ws.column_dimensions[letter].width = max(10, min(profile.get('widths', {}).get(letter) or default, 65))
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.print_title_rows = '1:1'
        ws.print_area = ws.dimensions
        ws.page_setup.orientation = profile.get('orientation') or ('landscape' if len(headers) > 5 else 'portrait')
        ws.page_setup.paperSize = profile.get('paperSize') or ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        if profile.get('margins'):
            ws.page_margins = copy(profile['margins'])
        ws.oddFooter.center.text = '第 &P 页 / 共 &N 页'
    sheet('项目概况', ['项目', '内容'], [['项目名称', project['title']], ['文书类型', project['kind']],
          ['申报单位', project['model'].get('school', '【待补充】')],
          ['编制状态', '待确认草稿' if project['issues'] else '待人工审核'],
          ['时间起点', project['model'].get('time_basis', '【待补充】')]])
    rows = []
    for section_data in sorted(project['sections'], key=lambda s: s['index']):
        for block in section_data['blocks']:
            for paragraph in block['paragraphs']:
                rows.append([section_data['title'], block.get('heading', ''), paragraph])
    sheet('章节内容', ['章节', '子标题', '正文'], rows)
    for table in shared_tables(project).values():
        sheet(table['title'], table['headers'], table['rows'])
    if include_notes:
        sheet('待确认事项', ['级别', '事项'], [[i['level'], i['message']] for i in project['issues']])
        citations = _citations_table(project)
        sheet('编制依据', citations['headers'], citations['rows'])
    wb.save(path)
    wb.close()


def export_project(project, out_dir, formats=('Word',), word_template=None, excel_template=None, include_notes=False, allow_draft=False):
    errors = [i['message'] for i in validate(project) if i['level'] == '错误']
    if errors and not allow_draft:
        raise ValueError('请先处理错误或继续完成正文：\n' + '\n'.join(errors))
    if errors:
        # 用户明知有错仍选择导出：强制附全部核对事项，只能是草稿
        include_notes = True
    directory = Path(out_dir)
    directory.mkdir(parents=True, exist_ok=True)
    result = {'paths': [], 'errors': []}
    stem = '%s_%s_%s_%s' % (_safe_name(project['title']), project['kind'],
                             datetime.now().strftime('%Y%m%d-%H%M%S'), uuid.uuid4().hex[:8])
    for format_name in formats:
        suffix = '.docx' if format_name == 'Word' else '.xlsx'
        path = directory / (stem + ('_待确认' if project['issues'] else '') + suffix)
        temporary = None
        try:
            fd, temporary = tempfile.mkstemp(dir=directory, suffix=suffix)
            os.close(fd)
            if format_name == 'Word':
                write_word(project, temporary, word_template, include_notes)
            elif format_name == 'Excel':
                write_excel(project, temporary, excel_template, include_notes)
            else:
                raise ValueError('未知输出格式。')
            # UUID冲突也不能覆盖已有文件。
            if path.exists():
                raise FileExistsError('目标文件已存在，请重新导出。')
            os.rename(temporary, path)
            result['paths'].append(str(path))
        except Exception as exc:
            result['errors'].append('%s保存失败：%s' % (format_name, exc))
        finally:
            if temporary and os.path.exists(temporary):
                os.unlink(temporary)
    return result
